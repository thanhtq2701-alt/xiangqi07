"""Build the editable V4 owner questionnaire with bundled python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'/'Co_Chill_V4_Tien_do_va_Bo_sung_Thong_tin.docx'
OUT.parent.mkdir(exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
 s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
 s.element.rPr.rFonts.set(qn('w:ascii'),'Arial');s.element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(28)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(7)
for s in ('Heading 1','Heading 2','Heading 3'):
 doc.styles[s].paragraph_format.keep_with_next=True

footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('Cờ Chill V4  |  08 tháng 10 năm 2026  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(110,110,110)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Cờ Chill V4 Tiến độ và thông tin cần bổ sung'
doc.core_properties.author='Cờ Chill · Trần Quốc Thanh · hỗ trợ biên soạn bởi Codex'
doc.core_properties.subject='Bản làm việc sản phẩm UI UX API và kế hoạch kiểm chứng'

def p(text='',bold=False,style=None):
 para=doc.add_paragraph(style=style);run=para.add_run(text);run.bold=bold
 return para

def h(text,level=1):doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title)
def bullet(text):p(text,style='List Bullet')
def answer(label='Trả lời',lines=1):
 para=p(label+' : ')
 para.paragraph_format.space_after=Pt(12*lines)
 para.runs[0].font.color.rgb=RGBColor(110,110,110)
 para.runs[0].italic=True
 return para

def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,title in enumerate(headers):t.rows[0].cells[i].text=title
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   if widths:cell.width=Inches(widths[ci])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'254B55' if ri==0 else ('F2F5F3' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2)
    for run in para.runs:
     run.font.size=Pt(9.3);run.bold=ri==0
     run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(25,35,35)
   cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p()

p('CỜ CHILL  |  BẢN LÀM VIỆC V4',True)
p('Tiến độ và thông tin\ncần bổ sung',style='Title')
p('Sản phẩm · Giao diện · Trải nghiệm học · Engine · API · Dữ liệu · Triển khai',style='Subtitle')
p('Dành cho Trần Quốc Thanh\nNgày cập nhật 08 tháng 10 năm 2026')
h('Bạn đang ở đâu',2)
p('Bạn đã có một prototype web chạy trên máy, chứng minh được vòng học từ ván cờ đến suy nghĩ, thử lại, bài liên quan và lịch ôn. Đây là nền tảng có thể kiểm tra và tiếp tục phát triển; chưa phải sản phẩm AI hoàn chỉnh hoặc website công khai.')
p('Bước tiếp theo có giá trị nhất là kiểm chứng chất lượng cờ và cách dạy bằng dữ liệu thật của bạn. Chưa cần mua API, gom hàng nghìn ván hoặc thêm PvP để chứng minh vòng học cơ bản.')
h('Tài liệu này giúp bạn làm gì',2)
bullet('Hiểu chính xác phần đã chạy, phần đang dùng giải pháp cục bộ và phần chưa được xây.')
bullet('Biết việc cần làm trong 60 phút tới, trong 3 ngày tới và trước khi mời người dùng ngoài thử.')
bullet('Trả lời 84 câu bổ sung có mã cố định, chia 14 nhóm; không cần trả lời lại toàn bộ V3.')
bullet('Gửi 3 thế cờ mẫu có lập luận chuyên môn để thay dữ liệu minh họa bằng nội dung đáng tin cậy.')
h('Cách trả lời ngắn nhất',2)
p('Giữ nguyên mã câu. Gõ “Giữ đề xuất”, “Chưa biết, Codex quyết”, hoặc quyết định khác kèm ví dụ. Với câu chuyên môn cờ, ưu tiên một thế cờ cụ thể hơn một đoạn mô tả chung. Bạn có thể trả lời tất cả, nhưng 12 câu ưu tiên được liệt kê ở phần hướng dẫn.')
p('Đây là tài liệu hỏi để chốt quyết định, không phải cam kết rằng mọi tính năng trong câu hỏi đã được triển khai.',True)

page('01 Tiến độ có bằng chứng')
p('Kết quả dựa trên mã nguồn hiện tại, kiểm thử tự động và thao tác trực tiếp trong trình duyệt. Không dùng phần trăm hoàn thành vì các hạng mục chưa có độ lớn tương đương.')
table(['Hạng mục','Trạng thái thật','Giới hạn còn lại'],[
('Bàn cờ và luật','Chạy được','Có luật 7 quân, tự chiếu, lộ Tướng, chiếu hết; chưa đủ phán quyết giải đấu.'),
('Chơi và nhập ván','Chạy được','Máy cục bộ, 2 người cùng máy, FEN và tọa độ; chưa PvP qua mạng.'),
('Xem lại ván','Chạy được','Tối đa 3 điểm từ 24 nước Đỏ gần nhất; dựa trên search nông.'),
('Tự suy nghĩ trước đáp án','Chạy được','3 câu bắt buộc trên server; chưa chấm được chất lượng lập luận.'),
('WHAT WHY HOW','Giải pháp thay thế','Mẫu tiếng Việt từ dữ kiện hợp lệ; không phải LLM hay HLV hiểu sâu.'),
('Thử lại và bài liên quan','Chạy được','12 thế tổng hợp, 4 kỹ năng; chưa được bạn/chuyên gia duyệt.'),
('Lịch ôn và tiến bộ','Chạy được','Quy tắc 1/2/7/14/30 ngày; điểm thực hành chưa chứng minh thành thạo.'),
('Lưu dữ liệu','Chạy cục bộ','SQLite và cookie khách; xuất/xóa dữ liệu; chưa tài khoản đa thiết bị.'),
('Kiểm thử','59 test lõi và 6 HTTP đạt','Có browser smoke toàn vòng học; chưa chứng minh 100 người đồng thời.'),
('Web công khai và AI mạnh','Chưa triển khai','Chưa hosting công khai, Pikafish, LLM, ML, OAuth, placement thật.')
],[1.35,1.35,3.9])
h('Mốc hiện tại',2)
p('Vòng học prototype đã hoạt động. Đang ở giai đoạn làm chắc nền và chuẩn bị kiểm chứng nội dung. Không nên gọi đây là “HLV AI tương đương chuyên gia”, “đã tăng trình người dùng” hoặc “đã train ML”.')
p('Browser smoke đã ghi 2 lần thử kiểm tra và 1 bài mới đúng, có lịch 14 ngày, còn sau tải lại. Đây là dữ liệu kiểm thử của Codex, không phải kết quả học của một người dùng thật.')

page('02 Bạn nên làm gì tiếp')
h('Trong 60 phút tới',2)
table(['Thời lượng','Việc làm','Đầu ra cần gửi'],[
('10 phút','Mở http://127.0.0.1:8765/ và đi hết demo.','Tên màn hình khiến bạn không biết bấm gì hoặc không hiểu ý.'),
('15 phút','Tự chơi 6–10 nước, kết thúc để xem lại.','Một nhận xét về bàn cờ, một nhận xét về cách dạy.'),
('20 phút','Chọn 3 thế cờ do bạn có quyền sử dụng.','FEN, nước đã đi, phương án chấp nhận, WHY/HOW và nguồn.'),
('15 phút','Trả lời 12 câu ưu tiên trong tài liệu.','Quyết định và ví dụ; ghi Chưa biết nếu chưa có bằng chứng.')
],[.85,2.8,2.95])
h('Trong 3 ngày tiếp theo',2)
bullet('Ngày 1: duyệt luật, sửa 3 lời giải mẫu theo cách bạn thật sự dạy; chốt tiêu chí “hiểu bài”.')
bullet('Ngày 2: chuẩn bị 10–20 thế cờ đa dạng và ghi nguồn; kiểm tra lời giải bằng engine mạnh hoặc chuyên gia độc lập. Chưa gọi tất cả là gold set khi chưa duyệt.')
bullet('Ngày 3: mời khoảng 3–5 người thử riêng nếu có thể; quan sát điểm mắc kẹt, không dẫn đáp án. Tách đánh giá dễ dùng khỏi đánh giá tiến bộ học tập.')
h('Phần Codex có thể tự làm',2)
p('Sửa lỗi, tăng độ rõ của trạng thái UI, giữ bản nháp, bổ sung điều hướng bàn phím, thiết kế API có phiên bản, adapter engine mạnh, schema dữ liệu và kiểm thử. Không cần bạn chọn framework hoặc tự viết code trước.')
h('Phần cần bạn quyết định hoặc cung cấp',2)
p('Lời giải chuyên môn đúng; phương án tốt nhưng khác engine; mức độ phù hợp người mới; dữ liệu có quyền dùng; phạm vi pilot; hạn mức chi phí thật; những đánh đổi trải nghiệm không có một đáp án khách quan.')
h('Chưa cần làm bây giờ',2)
p('Chưa mua domain/API chỉ để app “trông thật”. Chưa thu hàng loạt dữ liệu người dùng, xin mật khẩu hay gom puzzle từ sách/web chưa rõ quyền. Chưa thêm 3D, thanh toán hoặc PvP trước khi phần học đủ tin cậy.')

page('03 Kiến trúc và API hiện có')
p('API là cách giao diện trao đổi dữ liệu với chương trình xử lý. API nội bộ của app hiện miễn phí; chỉ khi tích hợp dịch vụ bên ngoài mới có điều kiện giá, quota và quyền truy cập riêng.')
p('Trình duyệt → HTTP JSON → Service → Rules và Engine → lời giải theo mẫu → SQLite → màn hình. Engine và lời giải tách riêng để thay thế sau này.')
table(['API hiện tại','Chức năng'],[
('GET /api/bootstrap','Cấu hình giao diện, kỹ năng và dashboard.'),
('POST /api/game','Tạo ván với máy, local, demo hoặc nhập FEN + moves.'),
('GET /api/game?id=…','Lấy ván thuộc phiên khách đang dùng.'),
('POST /api/move và /api/bot','Đi nước hợp lệ; máy chỉ đi khi đúng lượt.'),
('POST /api/review','Đóng ván để học và tạo các điểm xem lại.'),
('GET /api/exercise?id=…','Lấy bài; chưa phản hồi thì không trả lời giải.'),
('POST /api/reflect','Lưu 3 câu trả lời; review mới được mở phân tích.'),
('POST /api/attempt','Kiểm tra nước, lưu lần đầu và lên lịch ôn.'),
('POST /api/hint và /api/practice','Gợi ý có đánh dấu; chọn bài cùng kỹ năng chưa gặp.'),
('GET /api/dashboard','Tiến bộ, ván đã lưu, điểm xem lại và bài đến hạn.'),
('POST /api/feedback','Lưu góp ý cục bộ, không gửi ai.'),
('GET /api/export và POST /api/delete','Tải JSON hoặc xóa dữ liệu của phiên khách.')
],[2.45,4.15])
p('Hợp đồng hiện tại là prototype, chưa có OpenAPI/versioning, job queue, rate limit production hoặc tài khoản thật. Khi lên public phải bổ sung các phần này; không chỉ đổi URL localhost thành domain.')

page('04 Quyết định V3 đã được giữ')
p('Không cần trả lời lại các điều dưới đây, trừ khi bạn thật sự muốn thay đổi.')
bullet('Ngôn ngữ chính là tiếng Việt; trải nghiệm cờ tướng truyền thống, tập trung học và tư duy.')
bullet('Desktop trước, có bố cục thích ứng màn hình nhỏ; không để 3D/PvP chặn vòng học.')
bullet('Không mở đáp án trước tự suy nghĩ; sau đó có giải thích, thử lại và bài liên quan.')
bullet('Engine và luật quyết định dữ kiện cờ; LLM không được tự tạo nước đi, điểm hay biến.')
bullet('Không bắt buộc dịch vụ trả phí, đăng nhập production hoặc dataset lớn để chạy prototype.')
bullet('Không claim ML, người dùng, tác động học tập hoặc quyền dữ liệu khi chưa có bằng chứng.')
bullet('Ưu tiên bài cùng kỹ năng nhưng chưa gặp; tránh xem nhớ đáp án là hiểu bản chất.')
h('Các mặc định kỹ thuật đề xuất',2)
p('Giữ code đơn giản hiện tại đến khi yêu cầu thực tế buộc thay. API tương lai dùng phiên bản, validation rõ, quyền theo chủ bản ghi, idempotency cho attempt, timeout/cancel và job status cho phân tích dài. API key chỉ ở backend. Khi dịch vụ ngoài lỗi vẫn có đường học cục bộ được gắn nhãn.')
p('Mục tiêu UI tiếp theo: ít thao tác hơn nhưng không bỏ suy nghĩ; giữ bản nháp; lỗi có cách phục hồi; trạng thái máy đang nghĩ rõ; lời giải có ví dụ cụ thể; phân biệt tự giải và có trợ giúp; không tạo áp lực streak.')
h('Điều chưa thể tự suy ra từ V3',2)
p('Chúng ta chưa có bộ thế cờ chủ dự án duyệt, rubric lập luận, dữ liệu usability thật, độ khó đã hiệu chỉnh, thỏa thuận dùng dữ liệu bên ngoài, kế hoạch pilot cụ thể hay kết quả benchmark engine sâu. V4 tập trung lấp những khoảng trống đó.')
h('Nguyên tắc chốt sau khi bạn trả lời',2)
p('Một ý kiến sẽ được chuyển thành yêu cầu có điều kiện đạt: hành vi cụ thể, dữ liệu cần, giao diện liên quan và cách test. Ý tưởng chưa đủ bằng chứng sẽ nằm ở danh sách thử nghiệm; không tự biến tất cả thành tính năng phải làm ngay.')

page('05 Cách trả lời V4 hiệu quả')
p('Mỗi mã câu là một quyết định có ảnh hưởng đến sản phẩm. “Đề xuất” là phương án khởi đầu của Codex, không phải câu trả lời đã được bạn đồng ý.')
table(['Ký hiệu','Bạn làm gì'],[
('P0','Trả lời trước để chốt nội dung học và pilot gần nhất.'),
('P1','Trả lời sau khi thử giao diện, trước public alpha.'),
('Giữ đề xuất','Bạn chấp nhận phương án đang ghi.'),
('Chưa biết','Ghi điều bạn chưa biết; Codex sẽ chọn mặc định hoặc lập thử nghiệm.'),
('Bằng chứng','Kèm FEN, ảnh màn hình, lời người thử hoặc ví dụ cụ thể nếu có.')
],[1.25,5.35])
h('12 câu cần trước',2)
p('A01, A02, C03, E02, F01, F04, G02, J01, K01, L02, M01, N01.',True)
p('Đọc mục A đến N để trả lời đủ nếu bạn muốn. Không cần dành nhiều thời gian chọn màu nếu phần lời giải chuyên môn vẫn chưa được xác minh.')
h('Bản đồ câu hỏi',2)
table(['Nhóm','Phạm vi'],[
('A đến C','Mục tiêu pilot, thương hiệu/điều hướng, người mới và placement.'),
('D đến F','Bàn cờ, xem lại, tự suy nghĩ, WHAT WHY HOW và phản hồi.'),
('G đến I','Puzzle, bài học, lịch ôn, tiến bộ, mobile và khả năng tiếp cận.'),
('J đến L','Engine, API, LLM, chi phí, đăng nhập và dữ liệu cá nhân.'),
('M đến N','Nguồn dữ liệu, annotation, kiểm thử, triển khai và vận hành.')
],[1.25,5.35])
p('Khi gửi lại, giữ file có câu trả lời hoặc gửi text theo mã. Không điền API key, mật khẩu, token, thông tin nhận diện trẻ em hoặc dữ liệu chưa được phép chia sẻ vào tài liệu này.')

groups=[
('A','Mục tiêu và phạm vi kiểm chứng',[
('P0','Trong lần thử đầu, hành vi nào chứng minh app có ích: nêu được nước đáp, tự sửa sai, hay áp dụng ở thế mới? Cho một ví dụ cụ thể.','Chọn một mục tiêu chính là giải thích được ý tưởng và áp dụng vào một thế khác; chưa dùng tỷ lệ thắng làm chỉ số chính.'),
('P0','Ba đến năm người đầu tiên bạn có thể mời là ai theo nhóm trình độ, và họ đã biết luật đến mức nào?','Chỉ ghi bí danh và trình độ; bắt đầu nhóm nhỏ tương đối đồng đều để hiểu lỗi trải nghiệm.'),
('P1','Một phiên học thực tế nên kéo dài bao lâu, và điểm nào người học được phép dừng rồi tiếp tục?','Thử phiên 10–15 phút; hoàn tất một điểm xem lại và một bài liên quan; lưu giữa các bước.'),
('P1','Trong danh sách còn thiếu, xếp ba ưu tiên gần nhất và nói tính năng nào có thể chờ.','Độ đúng engine/nội dung, chất lượng lời giải, khả năng dùng trên thiết bị thật; PvP và 3D để sau.'),
('P1','Điều kiện nào khiến bạn cho người ngoài thử, và lỗi nào bắt buộc dừng phát hành?','Không còn lỗi luật/đáp án nghiêm trọng trong tập kiểm chứng; không mất dữ liệu; có đường báo lỗi và quay lại bản cũ.'),
('P1','Bạn muốn kết quả đợt này là demo nội bộ, alpha mời riêng hay public? Ngày mong muốn có phải hạn cứng không?','Chọn alpha mời riêng sau khi nội dung được duyệt; tách hạn demo khỏi hạn sản phẩm công khai.')]),
('B','Nhận diện và cấu trúc giao diện',[
('P1','Trong giao diện đang mở, ba chỗ nào khác cảm giác thương hiệu bạn muốn nhất? Ghi màn hình và thay đổi cụ thể.','Giữ bố cục học tập tối giản; đối chiếu xanh navy/vàng của V3 với xanh trầm hiện tại trước khi chốt bảng màu.'),
('P1','Logo Cờ Chill hiện có thuộc quyền của bạn không, và bạn có bản SVG/PNG gốc cùng quy tắc sử dụng không?','Dùng biểu tượng tạm do code tạo cho tới khi bạn cung cấp file có quyền dùng; không lấy logo từ trang khác.'),
('P1','Ở trang đầu, bạn muốn người mới bấm demo hay chơi thật trước? Trường hợp quay lại nên thấy gì đầu tiên?','Người mới thấy demo; người quay lại thấy bài đang dở hoặc kỹ năng đến hạn, không lặp onboarding.'),
('P1','Năm mục điều hướng hiện tại có dễ hiểu không? Mục nào nên đổi tên, gộp hoặc chuyển vị trí?','Giữ Chơi, Xem lại từ ván, Luyện, Tiến bộ và Trợ giúp dễ tìm; mỗi màn hình chỉ có một hành động chính.'),
('P1','Bạn muốn xưng hô và giọng phản hồi thế nào? Viết lại một câu app đang nói theo đúng giọng bạn.','Dùng bạn/mình hoặc câu trung tính, cụ thể, không phán xét; tránh gọi người học cẩu thả chỉ từ một nước sai.'),
('P1','Thông tin nào cần nằm ngay màn hình và thông tin nào chỉ cần mở rộng: nguồn, độ sâu, điểm, dòng biến, hạn chế?','Để WHY/HOW gần bàn cờ; kỹ thuật ở phần chi tiết nhưng vẫn có nhãn giới hạn engine dễ thấy.')]),
('C','Người mới và đánh giá đầu vào',[
('P1','Người chưa biết chữ Hán hoặc ký hiệu tọa độ cần trợ giúp gì để chọn đúng quân?','Có tùy chọn tên quân tiếng Việt, tooltip và bài hướng dẫn tương tác ngắn; không bắt xem video dài.'),
('P1','Dấu hiệu nào cho biết người học chưa biết luật và cần chuyển sang bài căn bản?','Dựa vào một vài tình huống luật, không suy ra từ tuổi hay chỉ từ tự khai trình độ.'),
('P0','Bạn hãy cung cấp hai thế placement, đáp án chấp nhận và dấu hiệu phân biệt người mới với intermediate.','Chọn hai kỹ năng khác nhau, ghi rõ cách xử lý người đúng nước nhưng không giải thích được.'),
('P1','Nếu tự khai giỏi nhưng placement thấp, app nên thông báo và cho thử lại thế nào?','Gọi đây là mức khởi đầu tạm; cho thêm bài xác minh, không gắn nhãn trình độ chắc chắn sau hai câu.'),
('P1','Mục tiêu cá nhân nào cần hỏi lúc đầu để thực sự thay đổi bài học, thay vì thu thông tin không dùng?','Hỏi tối đa một mục tiêu học có thể dùng ngay; phong cách tấn công/phòng thủ không được hợp thức hóa nước kém.'),
('P1','Lúc người mới không hiểu phải điền gì vào reflection, trợ giúp nào đủ mà không lộ đáp án?','Cho ví dụ cách mô tả trên một thế khác và nút Chưa biết ở từng câu; không điền sẵn lời giải thế đang làm.')]),
('D','Bàn cờ và trải nghiệm chơi',[
('P1','Sau khi tự chơi một ván, thao tác chọn quân và điểm đến có gây nhầm không? Cần kéo thả hay bấm hai lần đã đủ?','Giữ bấm chọn làm chuẩn cho chuột, cảm ứng và bàn phím; kéo thả chỉ là tùy chọn bổ sung.'),
('P1','Ván luyện cần cầm Đỏ/Đen, đổi bên và độ khó máy như thế nào để phù hợp mục tiêu học?','Bổ sung chọn bên sau khi review hai phía được kiểm thử; không đặt tên mức Elo cho bot chưa benchmark.'),
('P1','Khi bấm nước sai, bạn muốn app chỉ báo không hợp lệ hay giải thích lý do luật cụ thể?','Trong luyện luật giải thích rõ; trong ván chơi chỉ chỉ báo ngắn, không gợi nước hay hoặc đánh giá thế cờ.'),
('P1','Bạn cần đồng hồ và hình thức kết thúc ván nào ở alpha: đầu hàng, hòa, hết giờ, dừng để học?','Phân biệt ván luyện dừng để học với ván tính kết quả; chỉ thêm đồng hồ sau khi xử lý nền/reconnect đúng.'),
('P1','Việc lật bàn, lịch sử nước và ký hiệu nên hiện ra sao trên laptop và điện thoại?','Bàn cờ luôn là trọng tâm; history thu gọn trên mobile; chọn chuẩn notation trước khi hỗ trợ nhập nhiều loại.'),
('P1','Khi mất mạng hoặc máy tìm nước quá lâu, người chơi cần biết gì và làm gì tiếp?','Hiện trạng thái đang xử lý, timeout có nút thử lại; không gửi trùng nước; phục hồi ván từ server.')]),
('E','Xem lại và tự suy nghĩ',[
('P1','Một ván có nhiều lỗi thì nên học một lỗi lớn nhất hay đi qua tối đa ba lỗi theo thứ tự thời gian?','Mặc định một điểm quan trọng trước, cho xem thêm; ưu tiên cơ hội học rõ và bằng chứng chắc.'),
('P0','Ba câu reflection hiện tại có khai thác đúng cách bạn dạy không? Viết lại chúng cho một thế cụ thể.','Giữ đánh giá, nguy cơ, ứng viên và nước đáp; không thêm câu chỉ để làm người học chậm lại.'),
('P1','Với người giỏi, câu nào có thể gộp mà vẫn bảo đảm họ suy nghĩ trước khi xem lời giải?','Gộp đánh giá và nguy cơ; luôn cần ý tưởng/nước ứng viên cùng một lý do. Chưa cho bỏ hết reflection.'),
('P1','Khi rời màn hình, reload hoặc đóng trình duyệt giữa câu trả lời, bạn muốn giữ và phục hồi thế nào?','Tự lưu bản nháp riêng, phân biệt đã lưu nháp và đã gửi; không coi bản nháp là bài đã hoàn thành.'),
('P1','Người chơi tự nhận không biết nên nhận gợi ý ngay hay một câu hỏi dẫn đường nữa? Cho ví dụ.','Một gợi ý quá trình suy nghĩ trước; muốn xem đáp án vẫn được nhưng phải đánh dấu có trợ giúp.'),
('P1','Khi engine không chắc hoặc không tìm ra sai lầm rõ, màn hình review nên làm gì?','Nói chưa đủ căn cứ gắn lỗi; cho phân tích một điểm để học, không bịa blunder để ép đủ ba lỗi.')]),
('F','Chất lượng lời giải và cách dạy',[
('P0','Viết một lời giải WHAT WHY HOW theo đúng giọng HLV của bạn cho thế a2a7 trong demo hoặc một thế của bạn.','Nêu việc quan sát được, hệ quả qua nước đáp hợp lệ và một quy trình suy nghĩ có thể mang sang thế khác.'),
('P1','Một lời giải tối thiểu cần bao nhiêu biến, độ dài nào và yếu tố hình ảnh gì để người mới hiểu?','Một dòng chính ngắn, một đối chiếu với nước đã chơi và mũi tên theo lựa chọn; tránh đổ nhiều biến cùng lúc.'),
('P1','Nếu nước người học không trùng engine nhưng hợp lý, tiêu chí nào cho phép chấp nhận?','Dựa vào nhóm nước tương đương đã kiểm chứng hoặc ngưỡng cùng độ sâu; tách tốt khách quan khỏi dễ học.'),
('P0','Đưa rubric chấm lập luận gồm 3–5 tiêu chí và một câu trả lời đúng, một câu nửa đúng, một câu sai.','Đánh giá nhận diện đe dọa, ứng viên, nước đáp và kết luận; không dùng số từ hay từ khóa làm bằng chứng duy nhất.'),
('P1','Bạn muốn phản hồi thế nào khi người học đúng nước do đoán, hoặc sai nước nhưng nhận diện nguy cơ đúng?','Ghi nhận phần tư duy đúng và hỏi bước còn thiếu; không đồng nhất chọn đúng với hiểu bài.'),
('P1','Khi người dùng báo lời giải sai, quy trình nào xác minh, sửa và phản hồi cho họ?','Lưu phiên bản lời giải, FEN, engine/độ sâu và feedback; bạn duyệt mẫu trước khi sửa dữ liệu dùng chung.')]),
('G','Puzzle và thư viện bài học',[
('P1','Bốn chủ đề hiện tại thiếu chủ đề nào có giá trị nhất với nhóm pilot? Gửi một ví dụ mỗi chủ đề mới.','Thêm nội dung từ sai lầm thật trước, không tăng số lượng chỉ để kho trông lớn.'),
('P0','Hai thế như thế nào được coi là cùng ý tưởng nhưng đủ khác để kiểm tra chuyển giao? Đưa một cặp đạt và một cặp quá giống.','Khác bố cục, quân tham gia hoặc nước đáp; lật gương đơn thuần không mặc nhiên là bài hoàn toàn mới.'),
('P1','Một bài có nhiều nước đúng hoặc chuỗi bắt buộc nên được chấm thế nào?','Lưu tập đáp án/nhánh hợp lệ và điều kiện thành công; không ép một chuỗi tọa độ duy nhất nếu có phương án ngang giá.'),
('P1','Độ khó cần dựa vào yếu tố nào: độ sâu tính toán, số ứng viên, nhiễu hay kiến thức cần biết?','Dùng vài tiêu chí có thể kiểm tra, ghi mức tạm và hiệu chỉnh từ attempt thật sau này.'),
('P1','Khi hết bài chưa gặp cùng kỹ năng, bạn thích ôn bài cũ có nhãn, chuyển chủ đề hay chờ bổ sung?','Nói rõ đã hết bài mới; ưu tiên chủ đề liên quan và không ghi lặp lại là chuyển giao mới.'),
('P1','Một lesson nên kết thúc bằng điều gì để biết người học áp dụng được, và bạn có nội dung gốc nào dùng ngay?','Bài ngắn gắn một ví dụ và một thế kiểm tra; ghi quyền dùng cho cả văn bản, ảnh và vị trí cờ.')]),
('H','Lịch ôn và dashboard',[
('P1','Bạn muốn người học biết lịch ôn theo kỹ năng hay theo từng bài, và câu giải thích lịch nên nói gì?','Theo kỹ năng với bài mới nếu có; nêu vì sao ôn hôm nay, không tuyên bố khoảng cách là tối ưu khoa học.'),
('P1','Khi người học bỏ lỡ lịch vài ngày, app nên xử lý thế nào để không tạo áp lực?','Giữ hàng đợi nhỏ, không phạt điểm vì bận; ưu tiên điểm yếu và bằng chứng còn ít.'),
('P1','Điểm 0–100 đang hiển thị có hữu ích hay dễ bị hiểu nhầm là trình độ? Bạn muốn cách diễn đạt khác nào?','Dùng tiến độ thực hành và số bằng chứng; chưa dùng “thành thạo” khi chưa có lập luận tốt và chuyển giao đủ.'),
('P1','Bạn cần những tín hiệu nào để kết luận lỗi tư duy đang giảm trong ván thật?','Tách luyện puzzle khỏi thực chiến; đo tái diễn lỗi trong các cơ hội tương đương, không chỉ đếm lỗi thô.'),
('P1','Người học muốn nhắc ôn qua kênh nào và tần suất tối đa nào thực sự dễ chịu?','Bắt đầu nhắc trong app; email/push chỉ khi có đồng ý, dễ tắt và có nhu cầu đã xác minh.'),
('P1','Dashboard nên giúp người học quyết định hành động tiếp theo bằng thông tin nào?','Một ưu tiên luyện, bài đến hạn, bài đang dở và xu hướng có đủ dữ liệu; không dựng biểu đồ trống hoặc số liệu mẫu.')]),
('I','Mobile khả năng tiếp cận và lỗi giao diện',[
('P1','Người thử sẽ dùng model điện thoại, trình duyệt và kích thước màn hình nào? Có máy yếu hoặc mạng chậm không?','Ghi thiết bị thật thay vì “hỗ trợ mọi thiết bị”; ưu tiên kiểm thử các tổ hợp được dùng nhiều nhất.'),
('P1','Trên mobile, bàn cờ và câu hỏi nên cùng màn hình hay chuyển thành từng bước?','Bàn cờ vẫn xem lại được khi viết; tránh bàn phím che câu hỏi và tránh cuộn qua lại quá dài.'),
('P1','Có nhu cầu chữ lớn, tương phản cao, phân biệt quân không dựa vào màu, hoặc điều khiển bàn phím không?','Coi đây là yêu cầu nền; nhãn quân rõ, focus rõ, có phương án thay thao tác kéo và không dùng màu làm tín hiệu duy nhất.'),
('P1','Loading, lỗi mạng, dữ liệu trống, hết bài và hết quota cần nói gì để người học biết hành động tiếp theo?','Mỗi trạng thái có thông điệp cụ thể và nút phục hồi; giữ dữ liệu đã nhập và tránh chỉ hiện “Something went wrong”.'),
('P1','Âm thanh, chuyển động và rung có giúp ích hay gây xao nhãng trong giờ học của nhóm bạn?','Tắt mặc định hoặc rất nhẹ, có điều khiển; tôn trọng giảm chuyển động và không tự phát âm thanh dài.'),
('P1','Trong thử nghiệm UI, ba tác vụ nào bạn muốn người mới tự làm mà không được hướng dẫn?','Tự mở demo, tự thử lại, tự tìm lịch ôn; ghi thời điểm mắc kẹt và lời họ nói thay vì chỉ hỏi có thích không.')]),
('J','Engine và độ đúng chuyên môn',[
('P0','Chuẩn luật cho alpha cần theo tài liệu cụ thể nào, nhất là chiếu dai/đuổi dai/hòa? Bạn có văn bản hoặc ví dụ trọng tài không?','Giữ luật cơ bản hiện tại và gắn nhãn phần chưa đủ; không tự nhận tuân thủ luật Việt Nam khi chưa đối chiếu.'),
('P1','Bạn có engine hợp pháp đang dùng để đối chiếu không? Cho tên, phiên bản, cách chạy và phạm vi quyền sử dụng.','Ưu tiên adapter engine UCI cục bộ; chưa phân phối binary/weights trước khi kiểm tra điều khoản đúng phiên bản.'),
('P1','Độ chính xác và thời gian chờ nên đánh đổi thế nào khi xem lại 1 điểm so với toàn ván?','Phân tích nhanh có giới hạn rõ; nếu kết quả không ổn định thì nâng độ sâu ở điểm quan trọng, không đổi nhãn chắc chắn tùy tiện.'),
('P1','Cho một thế hy sinh quân mà đánh giá vật chất nông sẽ hiểu sai, cùng lời giải kiểm chứng.','Dùng làm regression case bắt buộc khi thay engine; không đánh giá chiến lược chỉ bằng số điểm quân.'),
('P1','Khi hai engine hoặc bạn và engine bất đồng, ai quyết và bằng chứng nào cần lưu?','Lưu cả kết quả/phiên bản/độ sâu, tìm sâu hơn và chờ domain review; không chọn ngầm kết quả đẹp hơn.'),
('P1','Bạn chấp nhận chỉ gắn nhãn critical khi các lần tìm kiếm ổn định và mức chênh đủ rõ không?','Định nghĩa ngưỡng theo độ sâu/loại kết quả; mate score tách centipawn; không dùng một con số cho mọi mức chắc chắn.')]),
('K','API LLM và ngân sách',[
('P0','Mốc dưới 100 nghìn đồng/tháng trong V3 còn giữ không, và bao gồm hosting hay chỉ LLM? Hạn chi cứng tối đa là bao nhiêu?','Giữ chế độ miễn phí cục bộ đến khi có nhu cầu rõ; không kích hoạt trả phí hoặc mua dịch vụ trước khi bạn duyệt ngân sách.'),
('P1','LLM cần cải thiện điều gì mà template chưa làm được: cách diễn đạt, hỏi tiếp, hay phản hồi lập luận? Xếp một ưu tiên.','Dùng engine facts làm đầu vào cố định; thử riêng phản hồi sư phạm, không giao LLM quyền tính nước cờ.'),
('P1','Bạn có tài khoản nhà cung cấp, gói API và hạn mức chi nào đã sẵn sàng không?','Chỉ ghi tên dịch vụ và trạng thái; không dán key/token. Thuê bao chat không tự động được giả định là có API quota.'),
('P1','Khi API lỗi, chậm hoặc hết ngân sách, trải nghiệm nào bạn muốn giữ bằng fallback?','Tiếp tục chơi, reflection, template, attempt và lịch ôn; nói rõ đang dùng giải pháp cục bộ, không retry vô hạn.'),
('P1','Có tích hợp bắt buộc nào ngoài engine/LLM như Google login, mail, phân tích hành vi hoặc import nguồn ván không?','Mỗi tích hợp phải có người dùng cần, dữ liệu gửi/nhận, quyền, giá và cách tắt; không thêm SDK chỉ để đủ danh sách.'),
('P1','Bạn muốn nhận báo cáo chi phí/lỗi ở đâu và ai được phép thay cấu hình model hoặc prompt?','Cấu hình backend có version, test trước phát hành, trần chi theo ngày/tháng; log bỏ dữ liệu nhạy cảm.')]),
('L','Tài khoản dữ liệu và quyền riêng tư',[
('P1','Lúc nào guest cần nâng lên tài khoản: lưu đa thiết bị, tham gia pilot hay ngay lần đầu?','Cho thử trước; chỉ yêu cầu tài khoản khi có lợi ích cụ thể. Không tự xây lưu mật khẩu nếu có provider phù hợp.'),
('P0','Pilot có người dưới 18 tuổi không, tại quốc gia nào, và ai chịu trách nhiệm rà soát thông báo/đồng ý trước khi thu dữ liệu?','Chưa mở thu dữ liệu trẻ em công khai khi chưa có quy trình được rà soát; không xem tài liệu này là tư vấn pháp lý.'),
('P1','Những dữ liệu nào thật sự cần cho học tập, và dữ liệu nào bạn đồng ý không thu?','Mặc định bỏ ngày sinh đầy đủ, trường/lớp, số điện thoại và tên thật; chỉ thêm khi có mục đích cần thiết đã rõ.'),
('P1','Cần giữ ván, reflection, logs và feedback bao lâu? Khi người học xóa tài khoản, backup xử lý thế nào?','Chốt lịch lưu theo loại dữ liệu và thời gian xóa; phân biệt xóa trong app với xóa các bản sao.'),
('P1','Nếu muốn dùng ván cho kho bài chung hoặc train model, người dùng được chọn độc lập và rút lại thế nào?','Tách đồng ý phục vụ học cá nhân khỏi chia sẻ/huấn luyện; dùng bí danh không được xem là đã ẩn danh hoàn toàn.'),
('P1','Ai được đọc reflection, sửa annotation hoặc tải dữ liệu khi có HLV/nhóm hỗ trợ?','Quyền tối thiểu, có nhật ký thay đổi; không gửi toàn bộ database cho người chỉ cần góp ý một thế cờ.')]),
('M','Nguồn dữ liệu và xác minh nội dung',[
('P0','Hãy gửi 3 thế đầu tiên theo mẫu cuối tài liệu, ưu tiên từ ván của bạn và những lỗi bạn từng thật sự dạy.','Cần FEN, lượt, nước đã chơi, phương án chấp nhận, dòng kiểm chứng, WHAT WHY HOW và tình trạng quyền.'),
('P1','Trong các sách/web/ghi chép đang có, nguồn nào được phép sao chép, sửa và phân phối? Có bằng chứng bằng văn bản không?','Lập sổ nguồn và phạm vi quyền; không suy ra quyền chỉ vì đọc được miễn phí hoặc tự nhập lại.'),
('P1','Bạn có thể duyệt bao nhiêu thế mỗi tuần mà vẫn kiểm tra được kỹ nội dung?','Chọn nhịp thực tế, duyệt theo checklist; 10 thế tốt có giá trị hơn 100 thế chưa chắc đáp án.'),
('P1','Annotation nào bắt buộc, và bạn muốn dùng Word, CSV hay công cụ biên tập để nhập?','Giữ schema chung gồm ID/FEN/side/solution/alternatives/skill/family/level/source/verification/version.'),
('P1','Một thế có lỗi sẽ bị gỡ, sửa phiên bản hay giữ cho người đã làm xem lịch sử?','Ngừng đề xuất bản lỗi; giữ lịch sử phiên bản, đánh dấu attempt bị ảnh hưởng và tính lại chỉ số khi cần.'),
('P1','Ai có thể kiểm tra độc lập một phần gold set để tránh mọi đáp án cùng phụ thuộc một người hoặc một engine?','Nhờ chuyên gia khi có thể; ít nhất ghi rõ chỗ chỉ tự duyệt và dùng tập kiểm tra giữ riêng.')]),
('N','Kiểm thử triển khai và vận hành',[
('P0','Kế hoạch thử thực tế gần nhất của bạn: ngày, số người dự kiến, thiết bị và nhiệm vụ họ làm?','Chạy thử riêng có quan sát; nếu chưa có người thì ghi chưa có, không thay bằng số người dùng giả.'),
('P1','Người thử cần truy cập ngoài máy của bạn ngay không, hay một buổi dùng chung máy đã đủ để kiểm tra cách học?','Nếu cần từ xa, làm public alpha có auth/quota/TLS/backup phù hợp; không mở thẳng dev server ra Internet.'),
('P1','Bạn có tài khoản hosting/domain nào đang sở hữu và có thể quản trị không? Ai giữ quyền billing?','Chỉ ghi dịch vụ và quyền truy cập sẵn có; chọn hosting sau khi đo tải engine và ngân sách, không chọn chỉ vì quen tên.'),
('P1','Lỗi nghiêm trọng cần được báo cho bạn bằng kênh nào, và trong bao lâu bạn có thể xử lý thực tế?','Có severity, reproduction, phiên bản và chủ xử lý; tắt tính năng lỗi hoặc quay về bản ổn định khi cần.'),
('P1','Bạn muốn đo thành công pilot ra sao mà không nhầm thích giao diện với học tốt hơn?','Đo hoàn tất tác vụ, hiểu lời giải và bài chuyển giao trễ; tách usability với learning, chưa kết luận nhân quả từ nhóm nhỏ.'),
('P1','Sau khi trả lời V4, bạn muốn ba tiêu chí nghiệm thu phiên bản kế tiếp là gì?','Chốt thành tác vụ quan sát được, bộ dữ liệu dùng test và ngưỡng đạt; viết rõ phần chưa làm để tránh scope tăng vô hạn.')]),
]

for code,title,questions in groups:
 page(f'{code} {title}')
 p('Trả lời theo trải nghiệm hoặc bằng chứng của bạn. Nếu đồng ý với đề xuất, chỉ cần ghi “Giữ đề xuất”.')
 for i,(priority,question,default) in enumerate(questions,1):
  para=p(f'{code}{i:02d}  {priority}  {question}',True)
  para.paragraph_format.keep_with_next=True
  para.paragraph_format.space_before=Pt(7)
  para=p('Đề xuất: '+default)
  para.paragraph_format.keep_with_next=True
  para.paragraph_format.space_after=Pt(3)
  for run in para.runs:run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(80,90,90)
  answer(lines=1)

page('06 Mẫu gửi một thế cờ đáng tin cậy')
p('Sao chép trang này cho mỗi thế. Bắt đầu với ba thế; chưa cần gom dataset lớn. Không gửi nội dung từ sách hoặc website nếu chưa có quyền phù hợp.')
for title,guide in [
('Mã và mục đích','Ví dụ THANH001. Kỹ năng muốn dạy và người học mục tiêu.'),
('Thế cờ và biên bản','FEN đầy đủ, bên đi; nếu không có FEN thì mô tả tọa độ từng quân hoặc đính kèm ảnh do bạn có quyền dùng.'),
('Nước đã chơi và ý định','Nước đã đi, người chơi lúc đó nghĩ gì; nếu không biết ý định thì ghi không biết, không suy đoán.'),
('Đáp án và phương án khác','Nước hoặc tập nước chấp nhận, biến kiểm chứng, nước đáp mạnh nhất đã xét.'),
('WHAT và WHY','Điều đã thay đổi; tại sao quyết định có vấn đề hoặc tốt, căn cứ bằng biến hợp lệ.'),
('HOW và kiểm tra hiểu bài','Quy trình suy nghĩ muốn dạy; câu hỏi theo dõi; tiêu chí nhận biết người học thật sự hiểu.'),
('Bài chuyển giao','Thế khác cùng kỹ năng, khác gì đủ để tránh nhớ đáp án; family để quản lý độ tương tự.'),
('Nguồn và xác minh','Nguồn, tác giả, quyền dùng, ai duyệt, engine/phiên bản/độ sâu, ngày kiểm tra và điều chưa chắc.')]:
 h(title,2);p(guide);answer()

page('07 Mẫu quan sát một buổi thử')
p('Mục đích là tìm vấn đề trong cách dùng và cách hiểu, không chứng minh sản phẩm tốt bằng lời khen. Không yêu cầu thông tin nhận diện không cần thiết.')
for title,guide in [
('Bối cảnh','Bí danh, trình độ tự khai, đã biết luật chưa, thiết bị/trình duyệt, ngày thử.'),
('Nhiệm vụ một','Tự vào demo và đến phần tự suy nghĩ. Họ mắc kẹt ở đâu? Có cần người hướng dẫn không?'),
('Nhiệm vụ hai','Sau khi xem lời giải, nhờ họ nói lại bằng lời mình và thử lại. Ghi đúng điều họ nói, không gợi ý đáp án.'),
('Nhiệm vụ ba','Cho một thế liên quan chưa thấy. Ghi nước, lý do, trợ giúp đã dùng và điều họ chuyển được từ bài trước.'),
('Theo dõi trễ','Nếu họ đồng ý, thử lại sau khoảng thời gian đã định bằng thế khác. Không hứa có tiến bộ chỉ từ một buổi.'),
('Phân loại vấn đề','Luật/đáp án, nội dung dạy, UI, tốc độ hoặc kỳ vọng sai. Ghi mức độ, cách tái hiện, ảnh nếu được phép.'),
('Quyết định tiếp theo','Một thay đổi đáng làm nhất; bằng chứng nào sẽ cho biết thay đổi có hiệu quả; chưa đủ dữ liệu thì ghi rõ.')]:
 h(title,2);p(guide);answer()

page('08 Đọc gì và tìm thông tin gì')
p('Đọc để giải quyết quyết định đang có, không học lan man toàn bộ AI. Mỗi nguồn dưới đây đi kèm một đầu ra nhỏ bạn có thể tự viết lại.')
table(['Ưu tiên','Tài liệu','Đọc xong làm được gì'],[
('Ngay','README.md và RUN_AND_DEMO_CODEX.md','Tự chạy server, làm demo, biết cách dừng và khởi động lại.'),
('Ngay','KNOWN_LIMITATIONS.md và AI_USAGE.md','Nói đúng phần nào là luật/search/template; không nhầm với ML.'),
('Tuần này','BEGINNER_LEARNING_GUIDE_CODEX.md','Theo được một nước đi từ UI qua API đến database.'),
('Tuần này','WXF Rules và tài liệu luật Việt Nam bạn chọn','Ghi những trường hợp chiếu dai/đuổi dai/hòa cần implement.'),
('Trước đổi engine','Official Pikafish repository','Hiểu UCI, version, cách benchmark và nghĩa vụ trước phân phối.'),
('Trước public','W3C WCAG và OWASP session guidance','Có checklist khả năng tiếp cận và kiểm tra phiên người dùng.'),
('Khi học API','MDN HTTP overview','Phân biệt request/response, GET/POST, status code và JSON.')
],[.9,2.6,3.1])
h('Nguồn chính thức đã đối chiếu',2)
for label,url in [
('WXF World Xiangqi Rules','https://www.wxf-xiangqi.org/images/wxf-rules/2018_World_XiangQi_Rules_English2018.pdf'),
('Pikafish chính thức','https://github.com/official-pikafish/Pikafish'),
('W3C WCAG 2.2 quick reference','https://www.w3.org/WAI/WCAG22/quickref/'),
('MDN HTTP overview','https://developer.mozilla.org/en-US/docs/Web/HTTP/Guides/Overview'),
('OWASP Session Management','https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html')]:
 para=p(label+'\n'+url)
 for run in para.runs:run.font.size=Pt(8.5)
p('Ngày tham chiếu 08/10/2026. Giá API, hạn mức, giấy phép của phiên bản tải về và yêu cầu pháp lý phải được kiểm tra lại trước khi chọn hoặc phát hành. Không cần gửi bí mật truy cập để trả lời V4.')

doc.save(OUT)
print(OUT)
print('Questions:',sum(len(g[2]) for g in groups))
print('Explicit pages:',1+len(doc.element.xpath('//w:br[@w:type="page"]')))
