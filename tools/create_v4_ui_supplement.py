"""Create separate editable V4 UI architecture and rules questionnaire."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/V4_BO_SUNG_UI_UX_KIEN_TRUC_LUAT.md'
OUT = ROOT / 'deliverables/Co_Chill_V4_Bo_sung_UI_UX_Kien_truc_Luat.docx'
PRIORITY = {1,2,3,6,7,8,10,11,12,15,16,17,19,21,24,27,28,31,32,34}


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.color.rgb = RGBColor(0,0,0)
        s.font.size = Pt({'Title':25,'Heading 1':17,'Heading 2':12,'Heading 3':11}.get(name,10.5))
        s.paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('V4 bổ sung UI UX kiến trúc luật  |  ').font.size = Pt(8)
    page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE'); footer._p.append(page)
    doc.core_properties.title = 'Câu hỏi bổ sung V4 về giao diện kiến trúc web và luật cờ tướng'
    doc.core_properties.author = 'Cờ Chill'
    doc.core_properties.subject = '272 câu hỏi có chỗ trả lời và mẫu bổ sung luật'
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    group, number, code_block, ids = None, 0, False, []
    for line in lines:
        if line.startswith('```'):
            code_block = not code_block
            continue
        if not line:
            continue
        if code_block:
            p = doc.add_paragraph()
            r = p.add_run(line); r.font.name = 'Consolas'; r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
            doc.add_paragraph('Bộ câu hỏi riêng để trả lời cùng V4', 'Subtitle')
            doc.add_paragraph('272 câu hỏi  |  34 nhóm  |  20 câu ưu tiên  |  08 tháng 10 năm 2026')
            continue
        if line.startswith('## '):
            title = line[3:]
            match = re.match(r'UX(\d+) ', title)
            group = int(match[1]) if match else None
            number = 0
            if group or title.startswith(('Mẫu phác','Mẫu gửi','Mẫu hợp','Cách đóng')):
                doc.add_page_break()
            doc.add_heading(title, 1)
            if title == 'Bản đồ các trang cần chốt':
                doc.add_paragraph('20 câu ưu tiên để định hướng bản thiết kế')
                doc.add_paragraph(', '.join(f'UX{g:02d}.01' for g in sorted(PRIORITY)))
                doc.add_paragraph('Tra cứu nhóm trong Navigation Pane của Word hoặc tìm mã câu bằng Ctrl F. Mỗi nhóm có tám câu; đánh dấu ưu tiên không có nghĩa các câu còn lại đã được đồng ý.')
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], 2)
            continue
        if group and line.startswith('- '):
            number += 1
            ident = f'UX{group:02d}.{number:02d}'
            ids.append(ident)
            heading = ident + ('  ƯU TIÊN' if group in PRIORITY and number == 1 else '')
            doc.add_heading(heading, 2)
            p = doc.add_paragraph(line[2:]); p.paragraph_format.keep_with_next = True
            p = doc.add_paragraph('Trả lời: ')
            p.paragraph_format.space_after = Pt(12)
            p.runs[0].italic = True
            continue
        p = doc.add_paragraph(line)
        if line.endswith(': …'):
            p.paragraph_format.space_after = Pt(13)
    assert len(ids) == 272 and len(set(ids)) == 272
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    # Validate the saved artifact, not just the authoring data.
    saved = Document(OUT)
    found = [p.text.split()[0] for p in saved.paragraphs if re.match(r'^UX\d{2}\.\d{2}(?: |$)', p.text)]
    assert found == ids
    assert sum(bool(re.match(r'^UX\d{2}\.\d{2}  ƯU TIÊN$', p.text)) for p in saved.paragraphs) == 20
    print(f'Created {OUT.name}: {len(found)} unique questions; 20 priorities; {len(saved.paragraphs)} paragraphs')


if __name__ == '__main__':
    build()
