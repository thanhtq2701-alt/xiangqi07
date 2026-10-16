"""Build the editable V5 questionnaire from its reviewed Markdown source.

Use the bundled Codex Python runtime with python-docx. The original owner
document is never opened or changed by this builder. Render the result before
delivery; tools/render_word_hidden.ps1 is a Windows fallback when the packaged
LibreOffice renderer is unavailable.
"""
from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'CO_CHILL_V5_QUESTIONS.md'
OUTPUT = ROOT / 'deliverables' / 'Co_Chill_V5_Cau_hoi_hoan_thien.docx'


def set_style(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.12
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is not None:
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            rf.attrib.pop(qn('w:' + attr), None)


def add_text(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    if text.startswith(('Đề xuất:', 'Trả lời:', 'Ưu tiên ')):
        label, _, rest = text.partition(':')
        if rest:
            p.add_run(label + ':').bold = True
            p.add_run(rest)
        else:
            p.add_run(text)
    else:
        p.add_run(text)
    p.paragraph_format.widow_control = True
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(27.94)
    sec.page_width = Cm(21.59)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)
    sec.footer_distance = Cm(.7)
    set_style(doc.styles['Normal'], 10.5)
    set_style(doc.styles['Title'], 25, True)
    set_style(doc.styles['Subtitle'], 13, False)
    set_style(doc.styles['Heading 1'], 16, True)
    set_style(doc.styles['Heading 2'], 12, True)
    set_style(doc.styles['Heading 3'], 11, True)
    for key in ['Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[key].paragraph_format.keep_with_next = True
        doc.styles[key].paragraph_format.space_before = Pt(10)
    set_style(doc.styles['List Bullet'], 10.5)
    # Remove template title rules, including inherited paragraph borders.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.core_properties.title = 'Cờ Chill V5 Câu hỏi hoàn thiện sản phẩm'
    doc.core_properties.subject = 'Bổ sung quyết định từ UI UX và Luật'
    doc.core_properties.author = 'Cờ Chill'
    doc.core_properties.keywords = 'V5, UI UX, luật, thử nghiệm, học cờ tướng'
    doc.core_properties.comments = ''

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Cờ Chill V5   |   ')
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    question_paragraphs = []
    in_question = False
    question_count = 0
    group_count = 0

    def finish_question():
        nonlocal question_paragraphs, in_question
        if not question_paragraphs:
            return
        # Keep a complete question and its response field together. Blank
        # writable paragraphs give room for short answers without table grids.
        answer = doc.add_paragraph('..........................................................................................................................')
        answer.paragraph_format.space_after = Pt(2)
        answer.paragraph_format.line_spacing = 1.0
        answer.runs[0].font.color.rgb = RGBColor(130, 130, 130)
        question_paragraphs.append(answer)
        blank = doc.add_paragraph('..........................................................................................................................')
        blank.paragraph_format.space_after = Pt(10)
        blank.paragraph_format.line_spacing = 1.0
        blank.runs[0].font.color.rgb = RGBColor(130, 130, 130)
        question_paragraphs.append(blank)
        for p in question_paragraphs[:-1]:
            p.paragraph_format.keep_with_next = True
        question_paragraphs[-1].paragraph_format.keep_with_next = False
        question_paragraphs = []
        in_question = False

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('# '):
            add_text(doc, line[2:], 'Title')
        elif line.startswith('## '):
            finish_question()
            heading = line[3:]
            if re.match(r'^[A-H] ', heading):
                group_count += 1
                p = add_text(doc, heading, 'Heading 1')
                p.paragraph_format.page_break_before = group_count == 1
            elif heading == 'Cách gửi lại V5':
                p = add_text(doc, heading, 'Heading 1')
            else:
                add_text(doc, heading, 'Subtitle')
        elif line.startswith('### V5 '):
            finish_question()
            question_count += 1
            in_question = True
            p = add_text(doc, line[4:], 'Heading 2')
            question_paragraphs.append(p)
        elif line.startswith('### '):
            finish_question()
            add_text(doc, line[4:], 'Heading 2')
        elif line.startswith('- '):
            add_text(doc, line[2:], 'List Bullet')
        else:
            p = add_text(doc, line)
            if line.startswith('Ưu tiên '):
                p.runs[0].italic = True
                for run in p.runs:
                    run.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(5)
            if in_question:
                question_paragraphs.append(p)
    finish_question()
    assert question_count == 41, question_count
    assert group_count == 8, group_count
    OUTPUT.parent.mkdir(exist_ok=True)
    doc.save(OUTPUT)
    print(f'Created {OUTPUT} with {question_count} questions in {group_count} groups')


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    build()
