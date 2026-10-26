from pathlib import Path
import hashlib
import json
import sys
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"E:\ThanhTQ\App cờ tướng\xiangqi-ai")
PACK = ROOT / "deliverables/Co_Chill_VSCode_2026_10_17"
WORK = ROOT / "var/reading-guide-2026-10-20"
OUT = ROOT / "deliverables/Co_Chill_Huong_dan_review_va_update_2026_10_20.docx"
WORK.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.66)
sec.left_margin = sec.right_margin = Inches(.72)
sec.header_distance = sec.footer_distance = Inches(.28)
for name in ["Normal", "Body Text", "List Bullet", "List Number"]:
    s = doc.styles[name]
    s.font.name = "Arial"
    s.font.size = Pt(11)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.07
for name, size in [("Title", 25), ("Heading 1", 19), ("Heading 2", 13)]:
    s = doc.styles[name]
    s.font.name = "Arial"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(9)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.keep_with_next = True
hp = sec.header.paragraphs[0]
hp.text = "CỜ CHILL   •   HƯỚNG DẪN REVIEW VÀ UPDATE"
hp.runs[0].font.size = Pt(8)
hp.runs[0].font.color.rgb = RGBColor.from_string("595959")
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run("Bản ngày 20/10/2026   •   Trang ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
fp._p.append(field)
doc.core_properties.title = "Hướng dẫn review và update Cờ Chill"
doc.core_properties.subject = "Tự đọc code, kiểm tra sản phẩm và quản lý đợt nâng cấp"
doc.core_properties.author = "Cờ Chill"

def p(text="", bold=False):
    x = doc.add_paragraph()
    r = x.add_run(text)
    r.bold = bold
    return x

def h(text):
    return doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")

def code(text):
    x = doc.add_paragraph()
    x.paragraph_format.space_before = Pt(3)
    x.paragraph_format.space_after = Pt(8)
    r = x.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    return x

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, row):
            c.text = str(text)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cannot = OxmlElement("w:cantSplit")
        trpr.append(cannot)
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + edge)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "D9D9D9")
                borders.append(e)
            pr.append(borders)
            mar = OxmlElement("w:tcMar")
            for edge in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "85")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            pr.append(mar)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "233D53" if ri == 0 else ("F1F5F8" if ri % 2 == 0 else "FFFFFF"))
            pr.append(shd)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.02
                for r in para.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
                    if ri == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

doc.add_heading("Hướng dẫn review và update Cờ Chill", 0)
p("Dành cho chủ sản phẩm mới bắt đầu đọc code", True)
p("Bạn cần kiểm tra sản phẩm theo thứ tự: yêu cầu đã chốt → thao tác trên web → code và dữ liệu → danh sách việc cần sửa. Sau đó chọn một đợt nâng cấp nhỏ có tiêu chí hoàn thành rõ ràng. Không cần thuộc toàn bộ code trước khi bắt đầu.")
p("Tài liệu này tiếp tục phần hướng dẫn đọc kiến trúc đang làm dở, đồng thời trả lời câu hỏi mới nhất: bây giờ review toàn bộ những gì đã có như thế nào và update tiếp ra sao.")
h("Việc làm ngay trong phiên review đầu tiên")
table(["Thứ tự", "Bạn thực hiện", "Sản phẩm cần có"], [
    ("1", "Xác định đúng bản web và chạy local theo trang 2.", "Ghi lại thư mục, địa chỉ và ngày kiểm tra."),
    ("2", "Đọc các tài liệu ưu tiên ở trang 3.", "Danh sách yêu cầu đã chốt và điểm còn chưa chốt."),
    ("3", "Thử vòng học và giao diện theo trang 7–8.", "Mỗi tình huống có kết quả Đạt, Lỗi hoặc Chưa thử."),
    ("4", "Ghi vấn đề theo mẫu ở trang 9.", "Bằng chứng, cách tái hiện và điều bạn muốn thay đổi."),
    ("5", "Chọn 3 việc cho đợt đầu ở trang 10.", "Phạm vi nhỏ để sửa, kiểm tra rồi mới phát hành."),
], [.6, 3.55, 2.85])
h("Bạn dùng hai tài liệu như thế nào")
p("Sổ tay 64 trang giải thích từng file là tài liệu tra cứu: mở đúng chương khi gặp file đó. File đang đọc là hướng dẫn thực hành: làm theo thứ tự, ghi kết quả và tự kiểm tra mức hiểu. Cả hai không thay thế việc trực tiếp chạy ứng dụng.")
p("Các trang sau có đường dẫn, bản đồ kiến trúc, lộ trình học, ca kiểm tra, mẫu ghi lỗi, kế hoạch nâng cấp và prompt bàn giao cho trợ lý AI. Trong Word, mở View → Navigation Pane để chuyển nhanh giữa các mục.")
p("Phạm vi kiểm tra ngày 20/10/2026: đọc tài liệu và mã nguồn local; chạy lại kiểm thử tự động của gói VS Code. Chưa kiểm tra thủ công lại toàn bộ giao diện hoặc xác minh website công khai trong lượt này. Các ca review bên dưới là việc cần làm, không phải kết quả đã đạt.", True)

page("1 Chọn đúng bản trước khi review")
p("Dự án có ba phạm vi. Chúng liên quan nhau nhưng không tự đồng bộ. Nếu mở nhầm bản, bạn có thể thấy giao diện hoặc cách lưu dữ liệu khác.")
table(["Phạm vi", "Vị trí và vai trò"], [
    ("Repository gốc", "E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai\nCó tài liệu lịch sử, source website và ứng dụng backend Python/SQLite."),
    ("Source website", "Thư mục online-preview trong repository gốc.\nHTML/CSS/JavaScript chạy trong trình duyệt; đây là nhánh trải nghiệm cần review trước."),
    ("Bộ source VS Code", "Thư mục deliverables\\Co_Chill_VSCode_2026_10_17 trong repository gốc.\nSnapshot độc lập gồm 29 file. Có 10 tài nguyên web cùng công cụ/test; không có backend đầy đủ hay lịch sử Git."),
], [1.45, 5.55])
p("Địa chỉ online đã được ghi trong báo cáo phát hành: https://co-chill-ngay-mai.cosmicanole32.chatgpt.site/. Báo cáo ghi phát hành phiên bản 3 ngày 16/10. Đây không phải xác nhận rằng nội dung đang phục vụ hôm nay vẫn khớp mọi file local.")
h("Mở bộ source để học trong VS Code")
p("Chọn File → Open Folder và dán đường dẫn dưới đây. Terminal phải đang đứng trong thư mục này, nơi có README.md và online-preview.")
code(r"E:\ThanhTQ\App cờ tướng\xiangqi-ai\deliverables\Co_Chill_VSCode_2026_10_17")
code("python --version\npython -m http.server 8769 --bind 127.0.0.1 --directory online-preview")
p("Mở http://127.0.0.1:8769/#home. Giữ terminal chạy; Ctrl+C để dừng. Nếu máy dùng Python Launcher, thay python bằng py. Nếu cổng đang dùng, kiểm tra server đã chạy trước khi mở thêm. Đổi cổng sẽ tạo một vùng lưu trình duyệt khác.")
p("Chạy website tĩnh không cần npm, API key hay database. Python ở đây chỉ đưa file đến trình duyệt. Ngược lại, lệnh python server.py --port 8765 tại repo gốc chạy ứng dụng Python/SQLite riêng.")
h("Chốt một nơi làm việc")
p("Học và thử sửa nhỏ trong gói VS Code được. Khi bắt đầu phát triển thật, chọn một checkout Git làm nơi sửa chính. Không sửa song song cả gói snapshot và repo gốc rồi đoán bản nào mới. Không xóa bản cũ hoặc tài liệu để làm repo gọn trong bước review.")

page("2 Đọc tài liệu theo thứ tự này")
p("Quy ước đường dẫn để bảng gọn: R = E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai. Khi thấy R\\docs, đó là thư mục docs bên trong đúng đường dẫn này. Đây là đường dẫn rút gọn trong tài liệu, không phải biến có thể dán vào terminal.")
table(["Ưu tiên", "File cần mở", "Bạn lấy thông tin gì"], [
    ("1", "R\\deliverables\\Co_Chill_VSCode_2026_10_17\\README.md", "Cách chạy đúng bản website, 10 file runtime, lệnh test và giới hạn."),
    ("2", "R\\BAO_CAO_UIUX_V5_2026_10_16.md", "Bản phát hành đã làm gì, chưa làm gì, lỗi đã ghi nhận và dữ liệu thật so với fallback."),
    ("3", "R\\docs\\UIUX_DECISION_REGISTER_2026_10_15.md\nR\\docs\\RULES_UIUX_2026_10_15.md", "Đối chiếu quyết định UI/UX và phạm vi luật. Dùng để xác định hành vi mong đợi."),
    ("4", "R\\deliverables\\Co_Chill_Giai_thich_tung_file_va_nhom_code_2026_10_18.docx", "Tra từng file, nhóm hàm và số dòng. Không cần đọc hết một lượt."),
    ("5", "R\\deliverables\\Co_Chill_V5_Cap_nhat_Tai_khoan.docx", "Bản 46 câu có bổ sung tài khoản. Dùng bản này thay bản V5 cũ 41 câu khi trả lời tiếp."),
    ("6", "R\\PROJECT_INVENTORY_VI.md\nR\\PROJECT_MEMORY_VI.md\nR\\GITHUB_CHECKLIST_VI.md", "Kho tài liệu, lịch sử quyết định và quy trình GitHub. Đối chiếu ngày vì có nội dung lịch sử."),
], [.65, 3.72, 2.63])
h("Đối chiếu ý định gốc khi có mâu thuẫn")
p("Tài liệu trả lời gốc nằm ở:\n• E:\\ThanhTQ\\App cờ tướng\\10-14\\V4.docx\n• E:\\ThanhTQ\\App cờ tướng\\10-14\\Thiết kế trang web.docx\n• E:\\ThanhTQ\\App cờ tướng\\10-13\\UI UX + Luat.docx")
p("Đặc tả nền: R\\docs\\spec\\00_START_HERE.md, 01_PRODUCT_SPEC.md, 03_ARCHITECTURE.md và 07_LEARNING_MODEL_SPACED_REPETITION.md. Answer V3.docx, Codex.docx và ZIP Final Pack vẫn ở R. Đó là nguồn yêu cầu và lịch sử; không có nghĩa mọi tính năng trong đó đã được xây.")
p("Khi thấy mô tả cũ khác báo cáo mới, ghi vào phiếu review cả hai nguồn và ngày. Ví dụ README gốc có phần backend cũ sau phần giới thiệu Site mới. Đừng lấy trạng thái của backend để kết luận website cũng đã có tính năng ấy.")

page("3 Hiểu kiến trúc bằng một thao tác")
p("Một trang web gồm phần người dùng nhìn thấy, phần quyết định hành vi và nơi lưu thông tin. Với website hiện tại, phần lớn công việc diễn ra ngay trong trình duyệt.")
code("Trình duyệt tải index.html\n  → style.css và các script theo thứ tự\n  → app.js nhận click hoặc form\n  → XQ kiểm tra luật   /   CC xử lý học và ván\n  → cập nhật state rồi save vào localStorage\n  → render dựng lại màn hình")
table(["Thành phần", "Hiểu đơn giản", "Trong Cờ Chill"], [
    ("HTML", "Khung chứa menu và nội dung.", "index.html có main; không cần mỗi trang một HTML."),
    ("CSS", "Màu sắc, cỡ chữ, khoảng cách và bố cục.", "style.css; đổi CSS không làm engine mạnh hơn."),
    ("JavaScript", "Phản ứng với thao tác và tính toán.", "app.js điều phối; các module còn lại cung cấp nghiệp vụ."),
    ("State", "Thông tin ứng dụng đang nhớ khi chạy.", "Ván, phiên học, lịch ôn và nháp suy nghĩ."),
    ("localStorage", "Kho nhỏ của trình duyệt tại địa chỉ này.", "Giữ trạng thái qua reload; không phải database chung của mọi người."),
    ("Backend và API", "Máy chủ xử lý yêu cầu và dữ liệu tập trung.", "Bản Site chưa có hệ thống tài khoản/API người chơi tập trung. Repo gốc có backend local riêng."),
], [1.15, 2.55, 3.3])
h("Thứ tự tải và ranh giới trách nhiệm")
p("index.html tải xiangqi.js → game-tools.js → fixtures.js → learning.js → experience.js → app.js. Các script dùng defer nên tải theo cơ chế của trang mà vẫn giữ thứ tự thực thi. XQ chứa luật/engine; CC được nhiều module bổ sung; FIXTURES chứa bài mẫu.")
p("Một URL như #home hoặc #review chỉ đổi phần giao diện do app.js dựng. Đây không phải đường dẫn đến một file home.html hay review.html riêng.")
h("Câu hỏi cần trả lời khi mở bất kỳ hàm nào")
p("Ai gọi hàm này? Nó nhận dữ liệu gì? Điều kiện nào chặn hoặc cho phép? Nó trả về gì? Nó có sửa state, ghi dữ liệu hay đổi giao diện không? File nào và test nào phụ thuộc vào nó?")
p("Chỉ khi trả lời được các câu này bằng một ví dụ cụ thể, bạn mới chuyển sang nhóm hàm tiếp theo. Không cần nhớ từng ký tự.")

page("4 Lộ trình đọc code không bị loạn")
p("Mỗi buổi chọn một mục nhỏ. Có thể dành 30–60 phút một buổi rồi dừng khi đạt tiêu chí. Đây là gợi ý phân chia, không phải cam kết số giờ để hiểu toàn bộ hệ thống.")
table(["Buổi", "Đọc và thực hành", "Dấu hiệu đã hiểu"], [
    ("1", "README, index.html, các vùng lớn của style.css. Chạy trang và xác định header, menu, main.", "Chỉ ra đâu là khung trang và đâu là nội dung được sinh sau."),
    ("2", "app.js: route, go, render và home. Bấm hai mục menu rồi theo dõi phần sau dấu #.", "Giải thích vì sao nhiều trang vẫn dùng một HTML."),
    ("3", "app.js: boardHTML, selectSquare, moveGame; xiangqi.js: legalMoves và play.", "Kể lại đường đi của một nước hợp lệ và một nước bị chặn."),
    ("4", "app.js: learn, finish, nextRelated; learning.js: createSession, canReveal, reveal, finishPuzzle.", "Phân biệt đang nghĩ, đã xem giải thích và đã hoàn thành."),
    ("5", "learning.js: schedule, recommend, sanitize; experience.js: reflectionCheck và safeExport.", "Nói được khi nào lịch ôn bắt đầu và dữ liệu nằm ở đâu."),
    ("6", "game-tools.js, fixtures.js và app/puzzles.py. Theo một bài hoặc một ván nhập.", "Phân biệt dữ liệu bài, luật đi quân và lời giải theo mẫu."),
    ("7", "tools/check_all.py, ba test JS và tests/test_rules.py. Chạy toàn bộ kiểm tra.", "Đọc được tên test lỗi và tìm module liên quan."),
    ("8", "AI_HANDOFF, SOURCE_MANIFEST, .gitignore và các tài liệu quản lý. Làm một sửa đổi nhỏ.", "Mô tả thay đổi, kiểm tra diff, thử lại và chuẩn bị một commit rõ mục đích."),
], [.55, 3.65, 2.8])
h("Cách đọc file dài hoặc nhiều code trên một dòng")
p("app.js có nhiều câu lệnh dài trên cùng dòng. Dùng Ctrl+F tìm tên hàm và Ctrl+G đến số dòng trong sổ tay. Số dòng chỉ là mốc của snapshot hiện tại; tìm theo tên hàm bền hơn sau khi code thay đổi.")
p("Nếu muốn format lại để dễ đọc, thực hiện trong một thay đổi riêng và kiểm tra toàn bộ sau đó. Không format hàng loạt rồi trộn sửa nghiệp vụ vào cùng lúc, vì sẽ khó biết phần nào gây lỗi.")
h("Sổ ghi chép sau mỗi buổi")
p("Ghi bốn dòng: hôm nay đọc file/hàm nào; một luồng mình đã hiểu; một câu còn chưa rõ; một bằng chứng đã tự kiểm tra. Câu còn chưa rõ phải cụ thể, chẳng hạn “finishPuzzle lấy thời điểm hoàn thành từ đâu?”")

page("5 Lần theo ba luồng quan trọng")
h("Luồng một từ bấm menu đến trang mới")
p("Bắt đầu từ index.html: tìm liên kết có hash. Sang app.js: tìm go, route và render. render chọn hàm nội dung như home, reviews hoặc profile rồi đưa HTML vào main. Bài tập: giải thích bằng lời vì sao thay text ở index.html có thể chỉ đổi menu, còn text trong trang phải sửa ở hàm trang tương ứng.")
p("Mốc snapshot app.js: go ở dòng 15, route ở 16, home ở 23, render ở 102. Các dòng chứa hàm rất dài; tìm tên đáng tin cậy hơn chỉ nhìn số dòng.")
h("Luồng hai từ chọn quân đến nước được lưu")
p("Trong app.js, vùng boardHTML tạo các ô có dữ liệu tọa độ. Sự kiện click chuyển ô sang selectSquare. Khi đủ điểm đi và đến, moveGame xử lý nước, gọi luật qua XQ và cập nhật ván. Luật trong xiangqi.js quyết định mục tiêu quân, an toàn Tướng và danh sách legalMoves.")
p("Mốc: app.js boardHTML dòng 24–35, moveGame dòng 83, selectSquare dòng 104–105; xiangqi.js targets dòng 92–147, legalMoves/play/status dòng 161–184.")
p("Bạn thử một nước Mã bị cản chân và một nước hợp lệ. Cần quan sát cả thông báo lẫn việc bàn cờ có thay đổi sai hay không. Muốn đổi màu thông báo thì sửa giao diện; muốn thay điều kiện nước hợp lệ thì phải sửa luật và thêm test tình huống.")
h("Luồng ba từ suy nghĩ đến lịch ôn")
p("app.js dựng ba ô suy nghĩ và nước dự kiến. Khi gửi, experience.js kiểm tra độ đầy đủ và quy tắc gợi mở; learning.js quản lý điều kiện mở lời giải. Tới khi hoàn thành, app.js gọi CC.finishPuzzle; learning.js ghi kết quả, xác định lịch ôn và dùng recommend để chọn bài tiếp.")
p("Mốc: app.js learn dòng 39–57, finish dòng 58, nextRelated dòng 59, form suy nghĩ dòng 151; learning.js canReveal/reveal dòng 129–148, schedule dòng 177–183, finishPuzzle dòng 185–238, recommend dòng 240–257.")
p("Điểm phải hiểu: mở hình chưa đồng nghĩa đã học xong. Lịch ôn lấy mốc hoàn thành hình; việc tự đánh giá cuối phiên không tự tăng kết quả khách quan. Trả lời đúng sau khi xem trợ giúp không được đánh đồng với đúng độc lập.")
h("Tự giải thích trước khi nhờ AI")
p("Vẽ bằng chữ: sự kiện → hàm → dữ liệu → điều kiện → kết quả → lưu ở đâu → test nào. Nếu thiếu một mắt xích, hỏi đúng mắt xích đó thay vì yêu cầu AI “giải thích cả dự án” lần nữa.")

page("6 Review vòng học cốt lõi")
p("Thử trên đúng bản đã ghi ở trang 2. Dùng dữ liệu thử riêng; không xóa tiến độ thật để tạo trạng thái sạch. Với mỗi ca, ghi Đạt, Lỗi hoặc Chưa thử và lưu bằng chứng khi cần. Cột mong đợi là tiêu chí kiểm tra, không phải kết quả đã được xác nhận hôm nay.")
table(["Mã", "Bạn làm gì", "Mong đợi cần kiểm tra"], [
    ("C01", "Vào trang chủ, bắt đầu một ván với máy hoặc hai người cùng máy.", "Bàn tải đủ quân, đúng bên đến lượt, không mở lời giải luyện tập trong lúc ván còn chạy."),
    ("C02", "Thử một nước sai luật, sau đó một nước hợp lệ.", "Nước sai không làm đổi trạng thái; nước hợp lệ đổi lượt và xuất hiện trong ký phổ."),
    ("C03", "Nhập một ván hợp lệ, rồi thử bản nhập có một nước sai.", "Ván hợp lệ vào thư viện; ván sai báo lỗi và không lưu một nửa ván."),
    ("C04", "Kết thúc ván rồi vào xem lại hoặc dùng bài demo.", "Có vị trí có căn cứ để học; không khẳng định mọi ván đều có sai lầm."),
    ("C05", "Gửi suy nghĩ thiếu một phần hoặc chưa chọn nước.", "Không lộ đáp án; chỉ rõ thông tin còn thiếu."),
    ("C06", "Gửi đủ suy nghĩ; thử một phiên cần gợi ý hoặc trả lời sai.", "Mở lời giải đúng quy trình; lưu câu đầu và trạng thái có hỗ trợ trung thực."),
    ("C07", "Đọc giải thích, xem biến minh họa và thử lại.", "WHAT, WHY, HOW liên quan đến đúng thế cờ; nước minh họa hợp lệ; thao tác thử lại chạy."),
    ("C08", "Hoàn thành hình và chuyển sang bài cùng kỹ năng.", "Bài tiếp ưu tiên chưa gặp, không tự nhận bài đã gặp là mới; hoàn thành chỉ được tính một lần."),
], [.6, 2.9, 3.5])
h("Ca đặc biệt cần thử kỹ")
p("Ở bước mời làm bài liên quan, chọn để lần sau rồi tải lại. Cần trở lại đúng chỗ đang dở. Sau đó hoàn thành bài và xem mốc lịch ôn. Đừng đổi đồng hồ hệ điều hành để thử; phần lịch nhiều ngày nên được kiểm tra bằng test thời gian có kiểm soát.")
p("Retry hiện là bước người học có thể thực hiện, chưa được cưỡng chế như điều kiện bắt buộc để hoàn thành. Nếu bạn muốn bắt buộc thử lại, ghi thành yêu cầu update riêng với tiêu chí rõ ràng.")
p("Thông tin ca kiểm tra: bản/thư mục ______  URL ______  ngày ______  thiết bị/trình duyệt ______")

page("7 Review giao diện dữ liệu và giới hạn")
table(["Mã", "Kịch bản", "Tiêu chí review"], [
    ("C09", "Đóng/mở lại tab ở cùng trình duyệt và cùng địa chỉ.", "Ván hoặc bài dở, nháp và tùy chọn cần được giữ trong phạm vi hiện hỗ trợ. Ghi lại mọi bước bị mất."),
    ("C10", "Hoàn thành phiên và mở tiến độ, lịch ôn.", "Chỉ tính hoạt động đã có; tự đánh giá không tăng thành tích; ít dữ liệu phải nói chưa đủ căn cứ."),
    ("C11", "Đổi theme, cỡ chữ, quân Hán/Việt, âm thanh.", "Đọc rõ, thao tác nhất quán, tùy chọn được giữ sau reload theo phạm vi tính năng."),
    ("C12", "Thử laptop và màn hình điện thoại khoảng 390 px.", "Không tràn ngang; bàn, nút chính và form không che nhau; người dùng biết bước tiếp."),
    ("C13", "Điều hướng bằng Tab và thao tác bàn bằng phím hỗ trợ.", "Thấy vị trí focus; không bị kẹt; nút và thông báo có tên dễ hiểu."),
    ("C14", "Xuất dữ liệu cá nhân trong Cài đặt.", "File đọc được, không lẫn đáp án nội bộ. Không giả định có phục hồi tự động từ file này."),
    ("C15", "Mở bằng trình duyệt/máy/địa chỉ khác.", "Hiểu rằng tiến độ hiện không đồng bộ; đây là giới hạn localStorage, chưa phải lỗi tài khoản."),
    ("C16", "Xem thư viện, nhập/xuất ký phổ, lật bàn và kết thúc ván.", "Ghi rõ lỗi hiển thị/khôi phục. Báo cáo cũ đã nêu lật bàn không giữ qua reload và đồng hồ khi kết thúc cần chỉnh."),
], [.6, 2.9, 3.5])
h("Cách góp ý giao diện để sửa được")
p("Đừng chỉ ghi “UI xấu”. Ghi trang nào, ở kích thước nào, phần nào gây khó, ảnh chụp hoặc bản phác, và hành vi mong muốn. Ví dụ: “Trang học trên điện thoại: sau khi chọn nước tôi không thấy nút gửi; muốn nút xuất hiện ngay dưới phần trả lời, không che bàn.” Đây là ví dụ cách viết, không phải lỗi vừa được xác nhận.")
h("Ba câu để review từng trang")
p("Người dùng vào trang để làm việc gì? Nút chính có giúp họ làm việc đó không? Sau khi bấm, họ có biết đã xảy ra điều gì và cần làm tiếp gì không? Áp dụng cho Home, Chơi, Xem lại, Học, Ôn tập, Tiến độ và Cài đặt.")
p("Bạn chưa thể xem số tài khoản đăng ký của mọi người trên Site hiện tại: chưa có tài khoản và dữ liệu tập trung. Đừng dùng số liệu trên trình duyệt của mình làm thống kê toàn hệ thống.")

page("8 Ghi kết quả thành danh sách có thể sửa")
p("Bạn có thể tạo một file REVIEW_BACKLOG.md bằng VS Code, một bảng tính hoặc điền trực tiếp vào bản sao Word này. Chỉ chọn một nơi quản lý chính; mỗi vấn đề có mã để tránh lặp và thất lạc.")
table(["Trường", "Cần ghi"], [
    ("Mã và loại", "Ví dụ R001; lỗi đang có, cải tiến UX, yêu cầu mới hoặc câu hỏi cần chốt."),
    ("Bản được kiểm tra", "Thư mục/URL, commit nếu có, ngày, trình duyệt, kích thước màn hình."),
    ("Cách tái hiện", "Các bước chính xác, dữ liệu đầu vào mẫu không chứa thông tin riêng."),
    ("Thực tế và mong đợi", "Hai câu riêng. Mong đợi dẫn tới quyết định UI/luật nào, hoặc là đề xuất mới của bạn."),
    ("Bằng chứng", "Ảnh/video ngắn hoặc lỗi console đã bỏ dữ liệu nhạy cảm. Không chỉ nói “không chạy”."),
    ("Ưu tiên và nghiệm thu", "P0/P1/P2, ai phụ trách, điều kiện để đánh dấu xong."),
    ("Trạng thái", "Chưa kiểm tra → Đã xác nhận → Đang sửa → Chờ review → Đã kiểm tra local → Đã phát hành."),
], [1.6, 5.4])
h("Ví dụ một yêu cầu cải tiến")
p("R001 — giữ hướng bàn sau tải lại. Nguồn: giới hạn đã ghi trong báo cáo phát hành 16/10; cần tái hiện lại trước khi sửa. Bước thử: mở ván → lật bàn → reload cùng URL. Mong đợi đề xuất: bàn giữ hướng vừa chọn. Ưu tiên P2 nếu không cản việc chơi. Nghiệm thu: giữ hướng qua reload, ván khác có hành vi được chốt, không ảnh hưởng lượt đi hay tọa độ.")
p("Đây là cách chuyển một ghi chú giới hạn thành đầu việc rõ ràng. Không được tự ghi “đã sửa” hoặc “đã kiểm tra” trước khi thực hiện.")
h("Phân biệt lỗi và thay đổi kỳ vọng")
p("Nước sai luật vẫn được đi là lỗi. Muốn bỏ một khoảng trống hoặc thêm animation thường là cải tiến. Muốn đăng nhập từ mọi máy là tính năng mới cần backend. Phân biệt ba loại giúp chọn đúng mức công việc, thời gian và người review.")
h("Cuối phiên bạn cần chốt")
bullet("Các ca đã thử và kết quả; ca chưa thử giữ nguyên nhãn Chưa thử.")
bullet("Ba vấn đề quan trọng nhất, mỗi vấn đề có điều kiện hoàn thành.")
bullet("Yêu cầu V5 còn thiếu thông tin, kèm một quyết định cần bạn trả lời.")
bullet("Không gộp một lỗi đã biết và mười mong muốn mới vào cùng một mục.")

page("9 Chọn đợt nâng cấp tiếp theo")
p("Chọn việc dựa trên tác động đến người dùng. P0 là không mở được web, mất/sai dữ liệu hoặc luật sai nghiêm trọng. P1 là vòng học bị kẹt hoặc kết quả/lịch ôn sai. P2 là cải thiện bố cục, cách diễn đạt và tiện ích chưa chặn việc học.")
h("Đợt một ổn định bản hiện có")
p("Chạy đủ C01–C16, xác nhận lỗi đã ghi và xử lý P0/P1 trước. Chốt rõ việc thử lại có bắt buộc không, cách tiếp tục bài dở, hành vi đồng hồ khi kết thúc và những số liệu nào được hiển thị. Nghiệm thu: đi trọn vòng chơi/xem lại → suy nghĩ → giải thích → thử lại → bài liên quan → lịch ôn → tiến độ mà không bị kẹt.")
h("Đợt hai hoàn thiện trải nghiệm")
p("Dựa trên ảnh và ghi chú của bạn, sửa từng màn hình ưu tiên: Home, Chơi, Học, Ôn tập, Tiến độ. Chốt hệ màu, chữ, khoảng cách, vị trí nút và cách thông báo lỗi theo các quyết định UI đã có. Kiểm tra cả laptop, điện thoại và bàn phím. Có thể chọn ba cải tiến nhỏ mỗi lượt; không cần viết lại toàn bộ website.")
h("Đợt ba tài khoản và quản trị V5")
p("Mục tiêu là người chơi đăng nhập từ máy nào cũng dùng được tài khoản của họ, và chủ sản phẩm xem được thống kê có quyền hạn. Phần này cần dịch vụ đăng nhập, API, database tập trung và phân quyền admin; chỉ đưa file lên hosting không tự tạo các thành phần ấy.")
bullet("Bạn chốt cách đăng nhập, thông tin bắt buộc, ai là admin, số liệu cần xem và cách người dùng xóa/xuất dữ liệu.")
bullet("Thiết kế dữ liệu: người dùng, ván, phiên học, lần thử, lịch ôn; phân biệt hoạt động ẩn danh với tài khoản.")
bullet("Làm một lát cắt hoàn chỉnh trước: đăng nhập → lưu một lần học → máy khác thấy đúng tiến độ → admin thấy thống kê được phép.")
bullet("Kiểm thử hai tài khoản A/B: A không đọc/sửa dữ liệu của B; tài khoản thường không mở API quản trị. Mật khẩu không lưu plaintext.")
p("Bản Python/SQLite cũ có thể cung cấp logic tham khảo, nhưng không thể coi là hệ thống cloud đã sẵn sàng. Cần đánh giá cụ thể phần tái sử dụng, bảo mật, sao lưu và chuyển dữ liệu; chưa chốt nhà cung cấp hoặc mua hạ tầng trong tài liệu này.")
h("Những việc chưa nên lấy làm điều kiện để bắt đầu")
p("Engine mạnh hơn, LLM, PvP online và bộ dữ liệu lớn nên có đầu việc riêng sau khi vòng học ổn định. Mọi đánh giá/nước đi vẫn phải do luật và engine xác định; lời giải không được tự bịa biến cờ. Thêm tài khoản cũng không biến thống kê học thành bằng chứng hiệu quả đã được kiểm chứng.")

page("10 Quy trình update một việc từ đầu đến cuối")
table(["Bước", "Bạn và người viết code làm gì"], [
    ("1 Chốt", "Chọn một mã backlog, viết điều kiện đạt và phạm vi không cần thay đổi trong lượt này."),
    ("2 Chuẩn bị", "Chọn repo chính, xem git status và diff hiện có. Ghi baseline trước sửa; không đè các thay đổi chưa rõ nguồn."),
    ("3 Sửa nhỏ", "Tìm đúng module; thêm test tái hiện cho lỗi logic quan trọng. Với thay đổi chữ/khoảng cách nhỏ, kiểm tra trực tiếp là đủ."),
    ("4 Kiểm tra", "Chạy test liên quan và toàn bộ kiểm tra của phiên bản đang sửa. Thử lại ca lỗi và vòng học có thể bị ảnh hưởng."),
    ("5 Review", "Đọc diff, xem ảnh trước/sau khi sửa UI; xác nhận số liệu và nhãn fallback vẫn trung thực."),
    ("6 Commit", "Lưu một thay đổi có mục đích rõ ràng; cập nhật tài liệu nếu hành vi/cách chạy thay đổi."),
    ("7 Phát hành", "Bạn duyệt bản cụ thể rồi mới push/deploy. Kiểm tra link sau phát hành và giữ bản trước để có thể quay lại."),
], [1.1, 5.9])
h("Lệnh kiểm tra tại bộ source VS Code")
code("python tools/check_all.py")
p("Lệnh yêu cầu Python 3.10+ và Node.js 18+ theo runner hiện có. Nếu node trên PATH quá cũ, truyền đường dẫn Node phù hợp. Phần DUONG_DAN_NODE bên dưới phải thay bằng vị trí thật, không dán nguyên.")
code('python tools/check_all.py --node "DUONG_DAN_NODE\\node.exe"')
h("Các thao tác Git cần hiểu")
code("git status --short\ngit diff\ngit diff --staged")
p("Các lệnh trên xem thay đổi, không xuất bản website. Sau khi đã chọn đúng file: git add theo tên file → xem lại git diff --staged → git commit với mô tả cụ thể. Tránh git add . khi repo còn nhiều tài liệu riêng, dữ liệu runtime hoặc file chưa rà soát.")
p("Commit lưu mốc trong Git local. Push đưa commit lên GitHub. Deploy đưa bản chạy lên hosting; một số dự án có thể cấu hình deploy tự động khi push. Phải kiểm tra kết nối đó trước khi push. Gói portable chưa có .git; đừng force-push một lịch sử mới đè repo cũ.")
p("Sau mỗi đợt, cập nhật README nếu đổi cách chạy, báo cáo/giới hạn nếu đổi hành vi, và nhật ký thay đổi. Tài liệu gốc V4/UI UX vẫn giữ nguyên; ghi quyết định mới vào tài liệu cập nhật có ngày.")

page("11 Kết quả kiểm thử và cách dùng AI")
h("Baseline đã chạy trong lượt này")
p("Ngày 20/10/2026, chạy tools/check_all.py trong gói Co_Chill_VSCode_2026_10_17 bằng Python 3.12.14 và Node v24.19.0. Kết quả 12/12 nhóm kiểm tra đạt, exit code 0, thời gian runner báo 25,7 giây.")
table(["Phần", "Kết quả"], [
    ("Cú pháp", "6 file JavaScript kiểm tra thành công."),
    ("Nghiệp vụ JavaScript", "19 test trạng thái học; 14 test công cụ ván; 18 test trải nghiệm. Tổng 51 test đạt."),
    ("Luật Python tham chiếu", "30 test đạt."),
    ("Bài tổng hợp", "12 bài đối chiếu thành công với nguồn và luật."),
    ("Đối chiếu JS và Python", "205 vị trí, 5.468 vị trí con, 25 ca FEN và 153 dòng engine được kiểm tra; kết quả PASS."),
], [2, 5])
p("Kết quả này chỉ áp dụng cho gói portable được chạy test. Không đồng nghĩa toàn bộ repository, giao diện trên mọi thiết bị, bảo mật hệ thống cloud hoặc bản online đã được kiểm thử. Các bước thủ công C01–C16 vẫn cần thực hiện. Không có sửa mã ứng dụng hoặc deploy trong công việc hướng dẫn này.")
h("Prompt để giao một việc cho Claude hoặc trợ lý khác")
p("Bạn có thể sao chép đoạn sau rồi thay các ô trong ngoặc vuông:")
p("“Đây là dự án Cờ Chill. Nơi sửa chính: [đường dẫn]. Hãy đọc README, AI_HANDOFF và quy tắc repo trước. Việc cần làm: [mã backlog và yêu cầu]. Bản hiện tại: [URL/commit]. Tái hiện: [các bước]. Mong đợi: [tiêu chí đạt]. Đọc code liên quan, đề xuất phạm vi rồi thực hiện sửa nhỏ trong phạm vi tôi giao. Dùng luật/engine làm căn cứ, không tự bịa nước cờ hoặc thành tích. Không xóa/giảm test để vượt lỗi. Báo lại file đổi, nguyên nhân, test và kết quả, giới hạn còn lại. Hỏi tôi trước push/deploy, đổi public, mua dịch vụ hoặc xóa dữ liệu. Không auto-approve.”")
h("Ba bài thực hành an toàn để bắt đầu")
bullet("Trong bản học riêng, tìm home trong app.js và sửa một câu hướng dẫn. Reload và chỉ ra đúng nơi thay đổi; kiểm tra không vỡ dấu nháy hoặc cú pháp.")
bullet("Trong style.css, chọn một khoảng cách của nút hoặc card và chỉnh nhẹ. So sánh laptop/điện thoại, focus và vùng bấm. Hoàn tác đúng dòng nếu xấu hơn.")
bullet("Đọc test lịch ôn trong test_preview_learning.cjs rồi giải thích mốc finish bằng lời. Sau đó chạy test không sửa dữ liệu tiến độ thật.")

page("12 Bản đồ mười file chạy website")
p("Tất cả file ở bảng này thuộc thư mục sau. Để mở một file, ghép tên file vào cuối đường dẫn hoặc chọn trong Explorer của VS Code.")
code(r"E:\ThanhTQ\App cờ tướng\xiangqi-ai\deliverables\Co_Chill_VSCode_2026_10_17\online-preview")
table(["File", "Vai trò và điểm bắt đầu", "Nếu thiếu"], [
    ("index.html", "Khung trang, main, menu, thứ tự script.", "Không có điểm mở/khởi tạo web đúng."),
    ("style.css", "Theme, bố cục, bàn cờ, responsive.", "Mất diện mạo và bố cục đúng."),
    ("app.js", "Trang, sự kiện, state/save/render. Bắt đầu ở route và render.", "Không có luồng tương tác ứng dụng."),
    ("xiangqi.js", "Luật FEN, nước hợp lệ, kiểm tra Tướng, engine nông.", "Chơi và phân tích không hoạt động đúng."),
    ("game-tools.js", "Ký phổ, nhập ván, lịch sử và quy tắc lặp trong phạm vi hỗ trợ.", "Hỏng các nghiệp vụ ván tương ứng."),
    ("fixtures.js", "12 bài tự dựng và dữ liệu lời giải.", "Không có bộ bài mẫu như hiện tại."),
    ("learning.js", "Phiên học, mở đáp án, hoàn thành, lịch ôn, đề xuất.", "Vòng học và tiến độ mất logic cốt lõi."),
    ("experience.js", "Tùy chọn, gợi mở, xuất dữ liệu, đồng hồ và tín hiệu tiến bộ.", "Hỏng khởi tạo/tính năng phụ thuộc."),
    ("logo-co-chill.png", "Hình logo có sẵn.", "Mất logo, không tự làm sai luật."),
    ("favicon.svg", "Biểu tượng tab.", "Mất biểu tượng tab."),
], [1.42, 3.52, 2.06])
p("Giữ đủ 10 file để mang nguyên trải nghiệm hiện tại sang thư mục mới. Copy code không copy tiến độ hoặc quyền quản trị hosting. Đây là các file runtime; các file ở trang tiếp theo hỗ trợ phát triển và vẫn nên giữ để sửa sản phẩm đáng tin cậy.")
p("Khi chỉnh bài tập, cần sửa nguồn bài và sinh lại fixtures có kiểm tra, không sửa tay lời giải rồi bỏ qua luật. Khi chỉnh luật, đối chiếu cả implementation JavaScript và tham chiếu Python theo các test.")

page("13 Bản đồ các file hỗ trợ phát triển")
p("Đường dẫn bên dưới tính từ thư mục gói VS Code đã ghi ở trang 2. Bảng bao quát 19 file còn lại; cùng 10 file trang trước tạo thành 29 file bàn giao, không tính cache sinh ra khi chạy test.")
table(["File hoặc nhóm file", "Vai trò và tác động nếu bỏ"], [
    ("app/__init__.py\napp/rules.py\napp/engine.py\napp/puzzles.py", "4 file Python tham chiếu luật, tìm kiếm và nguồn bài. Không phải backend của Site. Bỏ đi có thể vẫn mở trang, nhưng mất nguồn sinh/đối chiếu dữ liệu và các test phụ thuộc."),
    ("tests/__init__.py\ntests/test_rules.py", "2 file tổ chức và kiểm tra luật Python. Bỏ đi mất lớp kiểm tra luật quan trọng; không chứng minh web an toàn chỉ vì nó vẫn mở được."),
    ("tools/check_all.py\ntools/check_browser_rules.py\ntools/check_online_preview.py", "3 công cụ chạy toàn bộ kiểm tra, đối chiếu JS/Python và xác minh bài. Không phục vụ trang, nhưng cần cho quy trình sửa và review."),
    ("tools/export_preview_fixtures.py", "Sinh fixtures.js từ nguồn bài. Chỉ chạy khi có chủ đích cập nhật dữ liệu bài; sau đó xem diff và chạy test."),
    ("tools/test_preview_learning.cjs\ntools/test_preview_game_tools.cjs\ntools/test_preview_experience.cjs", "3 bộ kiểm thử JavaScript cho học, ván và trải nghiệm. Bỏ đi mất khả năng phát hiện lỗi tái phát tương ứng."),
    (".vscode/tasks.json\n.gitignore", "2 file cấu hình: shortcut chạy/test trong VS Code và quy tắc bỏ qua file khi Git theo dõi. Thiếu task vẫn dùng terminal; thiếu ignore tăng nguy cơ đưa dữ liệu không cần thiết vào commit."),
    ("README.md\nAI_HANDOFF.md", "2 tài liệu hướng dẫn chạy và bàn giao cho trợ lý khác. Không chạy trong trình duyệt, nhưng cần để người mới hiểu đúng phạm vi."),
    ("SOURCE_MANIFEST.json\nTHIRD_PARTY_NOTICES.md", "2 file truy vết nguồn/hash và ghi chú nguồn/quyền sử dụng. Không phải runtime; bỏ đi mất dấu vết và thông tin cần review."),
], [3.05, 3.95])
p("Các Word, ZIP nguồn, báo cáo và docs/spec ở repository gốc là tài liệu quản lý/sản phẩm, không phải file trình duyệt cần tải. Có thể tách nơi lưu tài liệu sau khi đã sao lưu và có quyết định, nhưng không xóa chỉ vì web không import chúng.")
p("Không đưa database, lịch sử người học, mật khẩu, token hoặc khóa API vào repo công khai. Gói portable hiện không chứa dữ liệu tài khoản hay thông tin truy cập hosting. .gitignore hỗ trợ lọc, không tự đảm bảo mọi bí mật đều được phát hiện.")

page("14 Khi nào bạn có thể tự update tiếp")
h("Tự kiểm tra bằng sáu câu")
table(["Câu hỏi", "Câu trả lời cần nắm"], [
    ("Muốn đổi màu, sửa đâu?", "Bắt đầu ở style.css. Thay đổi nghiệp vụ không thuộc CSS."),
    ("Muốn đổi hành vi của nút, tìm đâu?", "Từ markup/data-action trong app.js hoặc index.html tới event handler rồi tới hàm nghiệp vụ."),
    ("Ai quyết định nước hợp lệ?", "Luật trong xiangqi.js; lời giải không được tự bịa nước. Python là tham chiếu kiểm tra trong gói."),
    ("Copy code có mang tiến độ theo không?", "Không. localStorage thuộc trình duyệt/profile/origin; source và dữ liệu người học là hai thứ riêng."),
    ("Commit xong đã lên web chưa?", "Không tự động nếu chưa có pipeline. Cần hiểu remote, push và cấu hình deploy thực tế."),
    ("Muốn biết người chơi từ mọi máy?", "Cần tài khoản, backend/API, database và admin tập trung; dashboard local hiện không làm việc này."),
], [2.65, 4.35])
h("Tiêu chí đủ để làm một thay đổi nhỏ")
bullet("Bạn biết bản đang sửa nằm ở thư mục nào và có thể tự khởi động.")
bullet("Bạn lần được một thao tác từ giao diện tới hàm, dữ liệu và nơi lưu.")
bullet("Bạn mô tả được hành vi trước/sau, biết file nào nên đổi và file nào không liên quan.")
bullet("Bạn xem được diff và đọc kết quả test; biết phân biệt chưa thử với đã đạt.")
bullet("Bạn có thể quay lại bản trước bằng quy trình Git/bản sao phù hợp mà không xóa dữ liệu người dùng.")
h("Ba việc bạn nên gửi lại sau khi review")
p("1. Bảng C01–C16 với kết quả và ba lỗi/cải tiến quan trọng nhất.\n2. Câu trả lời V5 còn thiếu, đặc biệt đăng nhập, admin, dữ liệu và thống kê cần xem.\n3. Ảnh hoặc phác thảo các màn hình muốn đổi, kèm điều người dùng cần làm được sau thay đổi.")
p("Từ đó, đợt update kế tiếp sẽ có phạm vi cụ thể: việc nào sửa, tiêu chí đạt, module liên quan, cách test và phần cần bạn duyệt. Không cần viết lại tất cả code để bắt đầu.")
h("Phiếu ghi cho lần review đầu")
p("Bản và địa chỉ tôi đang xem: __________________________________\nBa mục ưu tiên: _____________________________________________\nĐiểm cần chốt với người viết code: _____________________________\nĐiều kiện để tôi đồng ý phát hành: ______________________________\nNgày review lại: _____________________________________________")
p("Lưu bản hướng dẫn này cùng các tài liệu bàn giao. Mỗi đợt cập nhật chỉ sửa trạng thái và quyết định mới trong sổ review; giữ ngày và nguồn để biết vì sao sản phẩm thay đổi.")

for part in [doc.styles.element, doc.element]:
    for node in part.xpath(".//w:pBdr"):
        node.getparent().remove(node)
doc.save(OUT)
source_hashes = {}
for f in sorted(PACK.rglob("*")):
    if f.is_file() and "__pycache__" not in f.parts and f.suffix != ".pyc":
        source_hashes[str(f.relative_to(PACK))] = hashlib.sha256(f.read_bytes()).hexdigest()
(WORK / "authoring_report.json").write_text(json.dumps({
    "output": str(OUT), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
    "manual_sections": 14, "source_hashes": source_hashes,
    "validation": "Portable check_all.py 12/12 PASS on 2026-10-20; no fresh live or full manual verification",
}, ensure_ascii=False, indent=2), encoding="utf-8")
print(str(OUT))
