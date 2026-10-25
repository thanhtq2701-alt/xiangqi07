from pathlib import Path
import hashlib
import json
import re
import sys
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[1]
PKG = ROOT / 'deliverables/Co_Chill_VSCode_2026_10_17'
OUT = ROOT / 'deliverables/Co_Chill_Giai_thich_tung_file_va_nhom_code_2026_10_18.docx'
snapshot = json.loads((WORK / 'snapshot.json').read_text(encoding='utf-8'))
inventory = {x['path']: x for x in snapshot}
items = []
for name in ['ui', 'logic', 'dev', 'support']:
    items += json.loads((WORK / (name + '.json')).read_text(encoding='utf-8'))
counts = Counter(x['path'] for x in items)
assert set(counts) == set(inventory), (set(inventory)-set(counts), set(counts)-set(inventory))
assert all(v == 1 for v in counts.values()), counts
for item in items:
    raw = (PKG / item['path']).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == inventory[item['path']]['sha256'], item['path']
    for group in item['groups']:
        if group['start'] is not None:
            assert 1 <= group['start'] <= group['end'] <= inventory[item['path']]['lines'], (item['path'], group)

order = [
    ('online-preview/index.html', 'Khung trang và thứ tự tải'),
    ('online-preview/style.css', 'Bố cục và giao diện'),
    ('online-preview/app.js', 'Điều phối toàn bộ trải nghiệm'),
    ('online-preview/xiangqi.js', 'Luật cờ và máy tìm nước'),
    ('online-preview/game-tools.js', 'Ký phổ và xử lý ván cờ'),
    ('online-preview/fixtures.js', 'Dữ liệu bài tập'),
    ('online-preview/learning.js', 'Phiên học và lịch ôn'),
    ('online-preview/experience.js', 'Tùy chọn và bằng chứng học'),
    ('online-preview/favicon.svg', 'Biểu tượng tab'),
    ('online-preview/logo-co-chill.png', 'Ảnh thương hiệu'),
    ('app/__init__.py', 'Nhận diện package Python'),
    ('app/rules.py', 'Luật tham chiếu bằng Python'),
    ('app/engine.py', 'Máy tìm nước tham chiếu'),
    ('app/puzzles.py', 'Nguồn bài tập tự dựng'),
    ('tests/__init__.py', 'Nhận diện package kiểm thử'),
    ('tests/test_rules.py', 'Kiểm thử luật Python'),
    ('tools/check_all.py', 'Điều phối bộ kiểm tra'),
    ('tools/check_browser_rules.py', 'Đối chiếu luật giữa hai ngôn ngữ'),
    ('tools/check_online_preview.py', 'Kiểm tra dữ liệu phát hành'),
    ('tools/export_preview_fixtures.py', 'Sinh dữ liệu bài tập cho browser'),
    ('tools/test_preview_game_tools.cjs', 'Kiểm thử nhập ván và ký phổ'),
    ('tools/test_preview_learning.cjs', 'Kiểm thử tiến độ học'),
    ('tools/test_preview_experience.cjs', 'Kiểm thử trải nghiệm và thống kê'),
    ('.vscode/tasks.json', 'Lệnh tiện ích trong VS Code'),
    ('.gitignore', 'Quy tắc bỏ qua của Git'),
    ('README.md', 'Hướng dẫn chạy và bảo trì'),
    ('AI_HANDOFF.md', 'Bàn giao bối cảnh cho trợ lý'),
    ('SOURCE_MANIFEST.json', 'Dấu vân tay bộ nguồn'),
    ('THIRD_PARTY_NOTICES.md', 'Nguồn tham khảo và giấy phép'),
]
by_path = {item['path']: item for item in items}
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.59); sec.page_height = Cm(27.94)
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(1.9); sec.right_margin = Cm(1.9)
sec.header_distance = Cm(.7); sec.footer_distance = Cm(.7)
for name, size, bold in [('Normal',11,False), ('Title',26,True), ('Subtitle',12,False), ('Heading 1',17,True), ('Heading 2',12,True), ('Heading 3',11,True), ('List Bullet',11,False)]:
    s = doc.styles[name]; s.font.name = 'Arial'; s.font.size = Pt(size); s.font.bold = bold; s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.line_spacing = 1.08; s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.widow_control = True
    for color in s.element.xpath('.//w:color'):
        for key in list(color.attrib):
            if 'theme' in key.lower(): del color.attrib[key]
    if name.startswith('Heading'):
        s.paragraph_format.keep_with_next = True; s.paragraph_format.space_before = Pt(10)
for element in doc.styles.element.xpath('.//w:pBdr'): element.getparent().remove(element)
doc.core_properties.title = 'Cờ Chill — Giải thích từng file và từng nhóm code'
doc.core_properties.subject = '29 file của bộ nguồn VS Code ngày 17/10/2026, phân tích ngày 18/10/2026'
doc.core_properties.author = 'Cờ Chill'
header = sec.header.paragraphs[0]
header.add_run('CỜ CHILL   SỔ TAY ĐỌC MÃ NGUỒN').font.size = Pt(8)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Bản đọc code 18/10/2026   •   ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def para(text='', style=None, keep=False):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.keep_with_next = keep
    return p

def heading(text, level=1, new=False):
    text = re.sub(r'[^\w\s]', ' ', text)
    text = re.sub(r' +', ' ', text)
    p = doc.add_paragraph(text, 'Heading ' + str(level))
    if new: p.paragraph_format.page_break_before = True
    return p

def labeled(label, text):
    p = para(); p.add_run(label + ' ').bold = True; p.add_run(text); return p

def code(text, keep=False):
    p = para(text, keep=keep)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
    return p

def bullet(text): return para(text, 'List Bullet')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    for i,w in enumerate(widths): t.columns[i].width = Cm(w)
    for i,value in enumerate(headers): t.rows[0].cells[i].text = str(value)
    for row in rows:
        cells = t.add_row().cells
        for i,value in enumerate(row): cells[i].text = str(value)
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+name); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        no = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            c.width = Cm(widths[ci]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'80'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'243F5B' if ri == 0 else ('F3F5F7' if ri % 2 == 0 else 'FFFFFF')); pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.03; p.paragraph_format.space_after=Pt(3)
                for r in p.runs:
                    r.font.name='Arial'; r.font.size=Pt(9.5); r.font.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    para()
    return t

def bookmark(p, name, identifier):
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(identifier)); start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(identifier))
    p._p.insert(0,start); p._p.append(end)

def anchor(text, name):
    p=para(); p.paragraph_format.space_after=Pt(4)
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),name)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'24554A'); props.append(color)
    under=OxmlElement('w:u'); under.set(qn('w:val'),'single'); props.append(under)
    run.append(props); t=OxmlElement('w:t'); t.text=text; run.append(t); link.append(run); p._p.append(link)

para('Hiểu từng file\nvà từng nhóm code\ncủa Cờ Chill', 'Title')
para('Sổ tay chi tiết dành cho chủ sản phẩm mới học phát triển web', 'Subtitle')
para('Phân tích ngày 18 tháng 10 năm 2026 • Bộ nguồn VS Code đóng gói ngày 17 tháng 10 năm 2026')
para('Tài liệu này đi từ cấu trúc website đến từng nhóm hàm, dữ liệu và kiểm thử trong 29 file của bộ nguồn bạn đang dùng. Mục tiêu là giúp bạn biết code nằm ở đâu, một thao tác đi qua những phần nào, sửa ở đâu và điều gì sẽ thiếu nếu bỏ một file.')
heading('Bộ code được giải thích', 2)
code(str(PKG))
para('Đây là đường dẫn gốc trong mọi ví dụ. Ghép đường dẫn tương đối của một chương với đường dẫn này để tìm đúng file. Mở cả folder bằng VS Code; dùng Ctrl+P để tìm tên file, Ctrl+G để đến dòng, Ctrl+F để tìm tên hàm. Trong Word, bật View > Navigation Pane để đi theo tiêu đề hoặc dùng các liên kết ở phần danh mục.')
heading('Phạm vi và độ tin cậy', 2)
para('Đã đối chiếu nội dung toàn bộ 29 file với dấu vân tay SHA256 tại thời điểm đọc. Số dòng trong tài liệu áp dụng cho bản này; khi định dạng hoặc sửa code, số dòng có thể đổi. Tên hàm và tên trường vẫn là cách tìm tốt hơn. Nhiều hàm app.js được viết gọn trên một dòng dài nên một dòng có thể chứa cả một quy trình.')
para('Tài liệu giải thích mã thực tế, không biến mục tiêu V5 thành tính năng đã hoàn thành. Bản web chạy trong trình duyệt, lưu dữ liệu cục bộ, dùng engine nông và bài tập tự dựng. Chưa có hệ thống tài khoản tập trung, đồng bộ đa máy, dashboard quản trị toàn bộ người chơi hoặc LLM gọi qua API.')
para('Trong lượt soạn tài liệu này không sửa mã ứng dụng, không cài công cụ, không commit, push hay deploy. Các file lịch sử, server.py và tài liệu đặc tả ở repo đầy đủ nằm ngoài bộ 29 file này; chúng không phải phụ thuộc để chạy bản Site tĩnh. Không nên xóa chúng chỉ vì chúng không có trong gói.')

heading('Bản đồ kiến trúc trước khi đọc code', new=True)
para('Website hiện tại có hai đường hoạt động riêng: người chơi dùng code JavaScript trong trình duyệt; người phát triển dùng Python và Node để kiểm tra hoặc tạo lại dữ liệu. Một chương trình HTTP tĩnh chỉ đưa file tới trình duyệt, không quản lý người chơi.')
code('HTTP server local hoặc hosting\n  └─ online-preview/index.html\n      ├─ style.css + logo-co-chill.png + favicon.svg\n      └─ script defer theo thứ tự\n          xiangqi.js       → XQ\n          game-tools.js    → CC\n          fixtures.js      → FIXTURES\n          learning.js      → bổ sung CC\n          experience.js    → bổ sung CC\n          app.js           → dùng các phần trên\n                             dựng HTML và xử lý thao tác\n                             đọc ghi localStorage\n\n+Đường phát triển riêng\n+  app/*.py + tests/*.py + tools/*\n+      → kiểm tra luật và quy trình học\n+      → tạo lại fixtures.js khi chủ động chạy exporter'.replace('\n+', '\n'))
para('XQ, CC và FIXTURES là các tên dùng chung giữa những script. Đây chưa phải kiến trúc React hoặc Next.js; không có bước npm build bắt buộc. Các module dùng lớp bọc để dùng được cả trong browser lẫn kiểm thử Node. app.js nối các thành phần đó với giao diện.')
table(['Lớp', 'Trách nhiệm', 'Không nên nhầm với'], [
    ('Hiển thị', 'HTML tạo khung; CSS tạo bố cục; app.js dựng màn hình và xử lý sự kiện.', 'Một file HTML riêng cho mỗi trang #.'),
    ('Luật và tìm kiếm', 'XQ kiểm tra nước hợp lệ, trạng thái ván và chọn nước bằng tìm kiếm nông.', 'LLM hoặc engine thi đấu chuyên nghiệp.'),
    ('Học và dữ liệu', 'CC quản lý bước học, lần thử, lịch ôn và dữ liệu đã lưu.', 'Database tập trung của tất cả người chơi.'),
    ('Bảo trì', 'Python tham chiếu, test Node và checker phát hiện lỗi trước khi phát hành.', 'Backend tự chạy cùng website.'),
], [3,8,6.2])

heading('Từ vựng và cách đọc một khối code', new=True)
table(['Từ hoặc ký hiệu', 'Nghĩa cụ thể trong bộ code này'], [
    ('Hàm / function', 'Một thao tác có tên. Gọi hàm với đầu vào; nó trả kết quả hoặc gây tác động như lưu dữ liệu, đổi màn hình.'),
    ('Tham số / return', 'Tham số là dữ liệu đầu vào; return là kết quả trả lại. Không phải mọi hàm đều sửa dữ liệu đã truyền vào.'),
    ('Object / array', 'Object gom trường có tên, ví dụ game.status; array là danh sách có thứ tự, ví dụ game.moves.'),
    ('Class / instance', 'Class mô tả một loại đối tượng. Position là khuôn cho các thế cờ; mỗi instance có bàn và lượt riêng.'),
    ('State', 'Ảnh chụp dữ liệu đang dùng: ván cờ, bài đang học, lịch ôn, hồ sơ cục bộ. Khác với HTML được vẽ ra từ dữ liệu đó.'),
    ('DOM / render', 'DOM là cây phần tử trang. Render dựng hoặc thay nội dung theo state; render không tự chứng minh dữ liệu đúng luật.'),
    ('Event / data-action', 'Sự kiện là click, submit, change… data-action gắn tên hành động vào nút để bộ xử lý biết người dùng muốn làm gì.'),
    ('FEN / move / ply', 'FEN là chuỗi mô tả thế cờ và lượt. Chuỗi nước dùng tọa độ nguồn–đích. Một ply là một nước của một bên, không phải đủ lượt hai bên.'),
    ('localStorage / JSON', 'localStorage giữ chuỗi theo khóa ở từng origin. JSON biểu diễn object thành văn bản để lưu, tải lại hoặc xuất.'),
    ('Route sau dấu #', 'Ví dụ #home hoặc #review chọn màn hình phía client. Không tạo request đến một backend route tương ứng.'),
    ('API nội bộ / API HTTP', 'XQ và CC cung cấp hàm cho code khác gọi. API HTTP là giao tiếp qua mạng với máy chủ; hệ tài khoản kiểu này chưa có.'),
    ('Test / assertion', 'Test dựng tình huống; assertion nêu kết quả phải đúng. Test đạt chỉ chứng minh các tình huống đã kiểm tra.'),
    ('Sanitize / validation', 'Lọc dữ liệu về dạng được chấp nhận và kiểm tra điều kiện. Không phải lớp xác thực người dùng hoặc mã hóa.'),
    ('Snapshot / hash', 'Snapshot là bản nguồn tại một mốc. Hash giúp nhận ra nội dung đổi; không chứng minh chương trình đúng hoặc đã deploy.'),
], [4,13.2])
para('Mỗi nhóm bên dưới có vị trí dòng, các tên nên tìm và lời giải thích theo mục đích. “Ví dụ” minh họa đường đi của dữ liệu; “Điểm cần giữ” chỉ điều kiện dễ bị phá khi sửa. Đây là phân tích theo nhóm ý nghĩa, không phải bản dịch từng dấu ngoặc hoặc bản sao hàng nghìn dòng code.')

heading('Danh mục từng file', new=True)
para('10 file đầu thuộc web chạy trong trình duyệt. 19 file sau hỗ trợ kiểm thử, phát triển và bàn giao. Các liên kết nhảy đến chương tương ứng; Word có thể yêu cầu giữ Ctrl khi nhấp.')
for i,(path,title) in enumerate(order,1):
    anchor(f'{i:02d}   {path} — {title}', 'file_' + str(i))
heading('Thứ tự học đề xuất', 2)
para('Lượt một: đọc 01 → 02 → 03 để hiểu trang xuất hiện và nhận thao tác. Lượt hai: đọc 04 → 05 → 07 → 08 để hiểu luật, ván và vòng học. Lượt ba: đọc 06 → 14 → 20 để hiểu nguồn bài; đọc các chương kiểm thử để biết cách kiểm chứng. Các file cấu hình ở cuối dùng tra cứu khi chạy, chuyển máy hoặc bàn giao.')

for i,(path,title) in enumerate(order,1):
    item=by_path[path]; meta=inventory[path]
    p=heading(f'{i:02d}  {title}', new=i in [1,3,4,5,6,7,8])
    bookmark(p,'file_'+str(i),i)
    code(path,keep=True)
    line_desc='ảnh nhị phân, không có dòng code' if meta['lines'] is None else f"{meta['lines']} dòng"
    para(f"Vị trí: {str(PKG / path)}",keep=True)
    para(f"Kích thước: {meta['bytes']:,} byte • {line_desc} • {len(item['groups'])} nhóm giải thích".replace(',', ' '),keep=True)
    labeled('Vai trò', item['role'])
    labeled('Liên hệ với', '; '.join(item['depends']) if item['depends'] else 'Không có phụ thuộc trực tiếp được khai báo ở file này.')
    labeled('Nếu thiếu file', item['missing'])
    if not item['groups']:
        para('File hiện trống. Không có hàm, biến hay khối thực thi để phân tích; vai trò đến từ tên và vị trí trong cấu trúc package.')
    for j,g in enumerate(item['groups'],1):
        heading(f'{i:02d}.{j:02d}  {g["title"].rstrip(".:")}',2)
        if g['start'] is None: locator='Tài nguyên nhị phân'
        elif g['start']==g['end']: locator=f"Dòng {g['start']}"
        else: locator=f"Dòng {g['start']}–{g['end']}"
        if g['symbols']: locator += '  |  ' + ', '.join(g['symbols'])
        code(locator,keep=True)
        para(g['explanation'])
        if g.get('example'): labeled('Ví dụ',g['example'])
        if g.get('risk'): labeled('Điểm cần giữ',g['risk'])
    if item.get('reading_tip'):
        heading('Cách đọc và sửa file này',2); para(item['reading_tip'])

heading('Theo một thao tác đi qua nhiều file',new=True)
heading('Ví dụ một nước cờ trong ván đang chơi',2)
para('Trong app.js, boardHTML() tạo các nút có data-square. Listener click ở dòng 107 lấy ô được bấm và gọi selectSquare() ở dòng 104. Lần bấm vào quân của bên đến lượt đặt selected; lần bấm ô đích ghép thành chuỗi nước. Hàm lấy luật từ XQ.Position và kiểm tra legalMoves(). Nếu sai, invalidMove() ở dòng 103 dùng CC.illegalMoveReason() để giải thích, không chuyển quân.')
para('Nếu hợp lệ, moveGame() ở dòng 83 gọi pos.play() để tạo thế kế tiếp. Nó cập nhật đồng hồ qua CC.applyClockMove(), thêm lịch sử, lấy trạng thái kết thúc và xét lặp qua CC.repetitionStatus(). Sau đó save() ghi state vào localStorage, render() dựng lại màn hình và maybeBot() gọi XQ.analyze() nếu đến lượt máy. CSS chỉ quyết định hình thức của bàn mới; CSS không quyết định nước đi được phép.')
code('click → selectSquare → legalMoves → moveGame\n      → play + clock + repetition → save → render → maybeBot')
heading('Ví dụ từ suy nghĩ tới lịch ôn',2)
para('Khi mở bài, beginPuzzle() tạo hoặc tiếp tục session, markSeen() chỉ ghi đã xem. Ba câu trả lời được listener input lưu dần. Chọn nước đi ở bàn ghi session.candidate sau khi luật xác nhận hợp lệ. Submit reflection-form ở dòng 151 gọi CC.reflectionCheck(): kiểm tra đã điền, đối chiếu nước với mẫu và đếm lần nộp. Phần chữ tự luận chưa được máy hiểu hoặc chấm chất lượng.')
para('Khi đủ điều kiện, CC.reveal() chuyển stage sang explain. Người học có thể bấm retry để thử lại tối đa ba lượt; selectSquare() ở dòng 105 xử lý lượt luyện này. Code hiện cho phép kết thúc từ màn giải thích nên không nên mô tả thử lại là bước bắt buộc trong mọi trường hợp. Một nước thử lại đúng không ghi đè session.candidate hoặc firstCandidate; lượt đầu và luyện lại có ý nghĩa khác nhau.')
para('Bấm finish gọi finish() ở dòng 58, rồi CC.finishPuzzle() ghi một attempt, thời điểm hoàn thành và dueAt. Mốc ôn là completedAt cộng khoảng ngày từ schedule(), không phải lúc mở bài. nextRelated() hỏi CC.recommend() chọn bài cùng kỹ năng; trạng thái related_pending giữ lời mời làm tiếp. Bản tự đánh giá cuối phiên được summarizeSession() ghi riêng, không dời lịch của các hình đã hoàn thành.')
code('beginPuzzle → phản hồi + nước hợp lệ → reflectionCheck\n            → reveal → thử lại nếu chọn → finishPuzzle\n            → schedule → recommend → tự đánh giá')
para('Hai ví dụ này cho thấy vì sao cần đọc cả nơi tạo giao diện và module xử lý dữ liệu. Một hàm pure có thể trả kết quả đúng, nhưng app vẫn phải gọi nó ở đúng thời điểm, giữ đúng state và lưu kết quả. Ngược lại, giao diện đẹp không thay thế được kiểm tra luật hoặc tính nguyên vẹn dữ liệu.')

heading('Khi muốn thay đổi sản phẩm thì bắt đầu ở đâu',new=True)
table(['Bạn muốn thay gì', 'Đọc và sửa trước', 'Kiểm tra ảnh hưởng'], [
    ('Màu và khoảng cách', 'style.css; class được tạo ở index.html và app.js.', 'Theme sáng tối, cỡ chữ, mobile, tương phản và trạng thái chọn.'),
    ('Thêm một nút', 'Hàm màn hình trong app.js, data-action và bộ nhận click.', 'Nút gọi đúng nhánh, xử lý lỗi, lưu state khi cần, không lộ đáp án trước phản hồi.'),
    ('Sửa luật quân', 'xiangqi.js và rules.py; test luật và đối chiếu.', 'Quân bị chặn, biên bàn, an toàn Tướng, nước cấm và tình huống kết thúc.'),
    ('Thêm bài tập', 'puzzles.py, exporter, fixtures.js sinh ra.', 'FEN hợp lệ, nước mẫu hợp lệ, skill/source đúng và không nhân bản ID.'),
    ('Đổi điều kiện hoàn thành', 'learning.js, app.js, experience.js và test học.', 'Phản hồi, thử lại, hoàn thành một lần, lưu tiếp tục và lịch tính từ hoàn thành.'),
    ('Đổi thống kê', 'experience.js và phần profile/progress trong app.js.', 'Định nghĩa mỗi số, thời điểm ghi, tách có hỗ trợ và tự làm; không gọi là điểm Elo/ML.'),
    ('Có tài khoản trên mọi máy', 'Cần thiết kế mới ranh giới auth, API, persistence và backend.', 'Bộ hiện tại không có dữ liệu tập trung. Không thể đạt chỉ bằng đổi localStorage key.'),
], [3.5,6.6,7.1])
heading('Những việc không tự xảy ra khi bạn sửa code',2)
bullet('Save chỉ lưu file trên máy. Website online chỉ đổi sau một lần phát hành thành công theo quy trình hosting.')
bullet('Commit tạo mốc Git; push gửi commit lên remote. Tùy cấu hình hosting mới có deploy tự động sau push; bộ code này không tự thiết lập điều đó.')
bullet('Copy source mang theo giao diện, logic và bài tập đóng gói. Nó không mang theo dữ liệu localStorage của trình duyệt hoặc quyền quản trị tên miền.')
bullet('Node chạy kiểm thử JavaScript. Python HTTP gửi file tĩnh. Cả hai không tự cung cấp đăng nhập và database cho người chơi.')

heading('Quy trình đọc thử và bảo trì bộ nguồn',new=True)
para('Bắt đầu bằng một thay đổi nhỏ có mục tiêu rõ. Trước khi sửa, mở đúng bộ nguồn và chạy được bản gốc. Đọc chương tương ứng, tìm tên hàm hoặc class trong VS Code, sau đó kiểm tra nơi gọi nó. Các tên xuất qua XQ/CC là hợp đồng giữa module nên không đổi tùy tiện.')
heading('Chạy web từ thư mục gốc của bộ nguồn',2)
code('python -m http.server 8769 --bind 127.0.0.1 --directory online-preview')
para('Mở http://127.0.0.1:8769/#home và giữ terminal đang chạy. Nếu Python của máy dùng launcher Windows, thay python bằng py. Cổng bận thì chọn cổng trống và nhớ rằng đổi cổng tạo một origin lưu dữ liệu khác. Đường dẫn terminal phải là thư mục Co_Chill_VSCode_2026_10_17.')
heading('Chạy kiểm tra sau thay đổi',2)
code('python tools/check_all.py')
para('Runner cần một Node hiện đại để kiểm thử JavaScript. Nếu node trong PATH quá cũ, dùng tùy chọn --node kèm đường dẫn Node phù hợp. Lệnh này chạy kiểm tra, không tự deploy và không tự sửa lỗi. Đọc toàn bộ lỗi đầu tiên, sửa nguyên nhân rồi chạy lại phần liên quan; không bỏ assertion để đạt màu xanh.')
heading('Kiểm tra bằng thao tác thật',2)
para('Mở trang chủ, đi một nước hợp lệ và thử một nước cấm. Kết thúc hoặc nhập ván rồi mở khoảnh khắc học. Trả lời trước khi xem giải thích, thử lại, hoàn thành và xem lịch ôn. Tải lại trang để kiểm tra lưu tiếp tục; thử theme, điện thoại và xuất dữ liệu. Dùng một hồ sơ browser thử riêng nếu không muốn ảnh hưởng dữ liệu học của bạn.')
heading('Cách dùng tài liệu khi chuyển sang AI khác',2)
para('Mở folder trong VS Code, cung cấp README.md và AI_HANDOFF.md, rồi yêu cầu trợ lý đọc các file liên quan trước khi sửa. Nêu một thay đổi cụ thể và điều kiện đạt. Yêu cầu nhận lại danh sách file đổi, lý do và kết quả kiểm tra. File Word giúp bạn hiểu và duyệt; mã chạy vẫn là các file trong folder, không phải nội dung chép từ Word.')
heading('Mốc và giới hạn của tài liệu',2)
para('Mọi vị trí dòng và số file trong sổ tay được kiểm tra với snapshot riêng ngày 18/10/2026. Việc lập tài liệu không phải một lần kiểm định lại toàn bộ luật thi đấu hoặc chất lượng engine. Các chương test giải thích mã test hiện có, không khẳng định mọi trường hợp ngoài tập test đều đúng. Báo cáo render và danh sách đối chiếu chỉ phục vụ kiểm tra tài liệu và được giữ ở thư mục làm việc, không đưa vào source chạy web.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
report={'output':str(OUT),'files':len(items),'groups':sum(len(x['groups']) for x in items),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'source_hashes_unchanged':True,'all_line_ranges_valid':True,'chapters':[{'number':i,'path':path,'groups':len(by_path[path]['groups'])} for i,(path,_) in enumerate(order,1)]}
(WORK/'build_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
