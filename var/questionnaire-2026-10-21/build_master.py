from pathlib import Path
import json, re, sys, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')
R=Path(r'E:\ThanhTQ\App cờ tướng\xiangqi-ai')
W=R/'var/questionnaire-2026-10-21'
OUT=R/'deliverables/Co_Chill_V5_Tong_hop_cau_hoi_phat_trien_2026_10_21.docx'
V=json.loads((W/'v5_audit.json').read_text(encoding='utf-8-sig'))
P=json.loads((W/'prior_answers.json').read_text(encoding='utf-8-sig'))
E=json.loads((W/'engine_accounts.json').read_text(encoding='utf-8-sig'))
md=(R/'docs/CO_CHILL_V5_QUESTIONS.md').read_text(encoding='utf-8')
old=[]
for m in re.finditer(r'^### V5 (\d+) ([^\n]+)\n(.*?)(?=^### V5 |^## |\Z)',md,re.M|re.S):
    num=int(m[1]); paras=[s.strip() for s in m[3].strip().split('\n\n') if s.strip()]
    a=next(x for x in V['questions'] if x['id']==f'V5-{num:02}')
    old.append(dict(id=f'V5 {num:02}',title=m[2].strip(),priority=a['priority'],meta=paras[0],body='\n'.join(x for x in paras[1:] if not x.startswith(('Đề xuất:','Trả lời:'))),proposal='\n'.join(x.removeprefix('Đề xuất:').strip() for x in paras if x.startswith('Đề xuất:')),fields=a['decision_fields'],source=f'V5 {num:02} giữ đầy đủ nội dung gốc',orig_id=a['id']))
assert [x['id'] for x in old]==[f'V5 {i:02}' for i in range(1,47)]

amend={
3:('Cần chốt việc hoàn thành hình có bắt buộc đi lại nước đã học sau lời giải không, hay chỉ được bỏ qua bước thử lại khi chuyển sang trạng thái học có hỗ trợ. Không coi bấm xem đáp án là đã tự giải đúng.', ['Bước thử lại bắt buộc để hoàn thành hay có ngoại lệ nào']),
4:('Bổ sung cho engine mạnh: ví dụ được chấp nhận phải cùng mục tiêu; khi Pikafish đổi kết luận hoặc chưa đủ ổn định, không dùng điểm máy như đáp án tuyệt đối. Ngưỡng điểm do kỹ thuật hiệu chuẩn bằng bộ thế đã duyệt.', ['Ví dụ nào bắt buộc phải công nhận nhiều đáp án']),
19:('Rà lại mức ưu tiên: nếu bạn muốn tự sửa câu hỏi/bài mà không nhờ người viết code, công cụ biên tập có cần đưa lên cùng bản V5 không?', ['Biên tập tối thiểu cùng V5 hay để sau']),
27:('Đối chiếu thêm mong muốn V4 nhắc trước hạn 30 phút: giữ kiểu nhắc trước hạn này hay chỉ nhắc các bài chưa làm theo giờ đã nêu?', ['Có giữ nhắc trước hạn 30 phút không']),
36:('Tách thời gian chấp nhận cho chọn quân, máy đi một nước, phân tích một hình và phân tích cả ván. Có thể ghi chưa biết để kỹ thuật đo rồi đưa hai phương án cho bạn duyệt.', ['Chờ máy đi một nước', 'Chờ một hình', 'Chờ cả ván']),
37:('Nhu cầu engine mạnh thêm chi phí CPU/máy chủ. Mức 100.000 mong muốn và trần 200.000 đồng/tháng là ràng buộc cũ, chưa tự tăng. Ghi rõ giữ hay điều chỉnh, tách ngân sách vận hành sản phẩm khỏi phí công cụ AI lập trình bạn dùng cá nhân.', ['Trần vận hành hiện tại', 'Phần dành cho engine, tài khoản/database và lời giải', 'Chi phí một lần/tên miền nếu có']),
39:('Đề xuất làm mốc Pikafish local trước là đề xuất của trợ lý, chưa được coi là bạn đã chọn. Hãy xếp thứ tự engine mạnh, tài khoản/admin, đồng bộ và UX trong phạm vi V5 đã chốt.', ['Có đồng ý một mốc Pikafish local trước bản tài khoản online không']),
40:('Hosting giao diện, API tài khoản/database và dịch vụ chạy engine là các thành phần khác nhau. Đưa giao diện lên Vercel không tự tạo Pikafish hoặc đồng bộ. Bạn chỉ cần nêu dịch vụ đang có và giới hạn; đội kỹ thuật chọn cách kết nối rồi báo để bạn duyệt.', []),
}
for q in old:
    n=int(q['id'][-2:])
    if n in amend:
        q['body']+='\n'+amend[n][0]
        q['fields']+=amend[n][1]

new=[]
pa_map={}
for a in P['missing_questions_not_duplicate_v5']:
    q=dict(title=a['title'],priority=a['priority'],body=a['question'],proposal=a['suggested_default'],fields=a['answer_fields'],source='; '.join(a['source_ids']),related=[f'V5 {int(x):02}' for x in a['related_v5']],orig_id=a['id'],group='J Giao diện và trải nghiệm từng bước',condition='Trả lời cho bản V5; nếu muốn để sau, ghi rõ phạm vi chưa làm.')
    if a['id']=='PA07':
        q['title']='Xin hòa khi chơi cùng máy hoặc chơi với bot'
        q['body']='Trong hai người một máy, lời mời hòa cần người còn lại xác nhận ở đâu, có nút từ chối và thời hạn chờ không? Với bot, cho xin hòa theo điều kiện nào và ai quyết định chấp nhận? Phần engine lỗi và đồng hồ khi sự cố trả lời ở câu dịch vụ engine ngừng hoạt động, không cần lặp.'
        q['fields']=['Xác nhận hoặc từ chối hòa local','Thời hạn chờ và đồng hồ trong lúc xin hòa','Điều kiện hòa với bot','Cách ghi kết quả']
    if a['id']=='PA09':
        eng=next(x for x in E['questions'] if x['id']=='N-ENG09')
        q['body']+=' '+eng['question']
        q['fields']+=['Cảnh báo thiếu lịch sử','Cách bổ sung chuỗi nước hoặc ghi Không còn biên bản']
    if a['id']=='PA13':
        eng=next(x for x in E['questions'] if x['id']=='N-ENG07')
        q['body']+=' '+eng['question']
        q['fields']+=['Có xem quyết định tốt khi có bằng chứng không','Số vị trí có được ít hơn ba không']
    if a['id']=='PA17':
        q['title']='Công cụ tự suy nghĩ trước khi mở lời giải'
        q['body']='Trước khi gửi suy nghĩ, người học có được tự vẽ mũi tên, đánh dấu quân hoặc ghi nhiều nước ứng viên mà chưa nhận phản hồi máy không? Nếu tự mô phỏng vài nước hợp lệ thì lần đó có còn được gọi là tự giải không? Phần khám phá sau lời giải trả lời riêng ở nhóm Engine.'
        q['fields']=['Công cụ được dùng trước khi gửi','Có tự mô phỏng nước trước đáp án không','Cách phân biệt tự thử với nhận trợ giúp','Cách xóa dấu và trở về hình gốc']
    new.append(q)

extras=[
('Hình thức tổng thể cần giữ và cần thay','Bạn đã chọn ba theme và bố cục chính. Để hoàn thiện chi tiết, hãy gửi một ảnh web hiện tại có đánh dấu phần giữ, phần đổi; chọn cảm giác cho nút/card (phẳng hay có bóng nhẹ), mức bo góc và độ dày đường viền theo ví dụ trực quan. Nêu tối đa ba điều khiến giao diện trông thiếu chỉn chu; màu thương hiệu và logo vẫn trả lời V5 20.', ['Ảnh đánh dấu giữ và đổi','Nút, card, bo góc, bóng và viền mong muốn','Mức chữ và độ thoáng dễ đọc','Ví dụ chỉ tham khảo bố cục, không sao chép tài sản'], 'Kỹ thuật làm một mẫu Home và một mẫu màn học với đủ ba theme để bạn duyệt; dùng cùng hệ chữ/khoảng cách cho cả web, không yêu cầu bạn chọn từng pixel.', [20,21,22], 2),
('Chi tiết bàn cờ và trạng thái quân cần nghiệm thu','Giữ kiểu quân Hán/Việt và thao tác bấm đã chọn. Trên từng theme, trạng thái chọn quân, nước trước, chiếu và nước bị cấm có điểm nào khó nhận ra? Có cần kéo thả như phương án phụ, đánh số tọa độ thường trực, hoặc điều chỉnh tốc độ quân chạy khi xem biến không?', ['Ảnh trạng thái khó nhìn','Kéo thả có cần ở đợt này không','Tọa độ hiện thường trực hay bật tắt','Tốc độ chuyển nước và chế độ giảm chuyển động'], 'Bấm quân/đích và bàn phím luôn dùng được; không thêm chấm gợi nước hoặc đánh giá chiến thuật trong ván. Chọn hình và độ tương phản dựa trên thử thực tế.', [21,23,24], 2),
('Các màn tài khoản cần tên và đường đi nào','Từ Home, đăng ký/đăng nhập, xác minh, quên quyền truy cập, hồ sơ và đăng xuất cần nối nhau thế nào? Sau đăng nhập thành công hoặc hết phiên giữa bài, quay về đúng bài hay về Home? Trang hồ sơ cần trường nào được sửa và trường nào chỉ xem? Phương thức đăng nhập chốt tại V5 32.', ['Lối mở đăng nhập và câu mời','Trang quay về sau đăng nhập/hết phiên','Các trường hồ sơ được xem/sửa','Cảnh báo khi dùng máy chung hoặc đăng xuất có nháp'], 'Quay lại việc đang làm sau xác thực khi còn phù hợp; báo phần nháp chưa đồng bộ trước đăng xuất, không yêu cầu tên thật nếu không có mục đích.', [32,44,45,46], 1),
('Bố cục quản trị từ tổng quan tới chi tiết','V5 28 và 42 đã hỏi số liệu và quyền xem. Với các quyền đó, bạn muốn bắt đầu ngày làm việc ở tổng quan, danh sách tài khoản hay phản hồi lỗi? Khi bấm một tài khoản, cần các tab nào và hành động quan trọng đặt đâu? Bạn có cần dùng admin trên điện thoại không?', ['Trang mặc định và thứ tự menu admin','Thông tin trong hàng danh sách tài khoản','Các tab chi tiết theo quyền được phép','Hành động khóa/xuất/xóa và bước xác nhận','Có quản trị trên điện thoại không'], 'Tổng quan dẫn tới danh sách có lọc, rồi chi tiết được phân quyền; thao tác phá hủy tách khỏi nút xem thường dùng, không hiển thị câu trả lời riêng tư ngoài mục đích đã thông báo.', [28,38,42,43,46], 1),
]
for i,(title,body,fields,proposal,related,priority) in enumerate(extras,1):
    new.append(dict(title=title,body=body,fields=fields,proposal=proposal,priority=priority,related=[f'V5 {x:02}' for x in related],source='Rà soát bổ sung ngày 21/10/2026 từ UI UX và phạm vi tài khoản',orig_id=f'VIS{i:02}',group='K Hoàn thiện hình thức và các màn tài khoản',condition='Trả lời; được ghi giao kỹ thuật đề xuất mẫu để duyệt.'))
for a in E['questions']:
    if a['id'] in ['N-ENG07','N-ENG09']: continue
    group={'Engine':'L Pikafish và phân tích có căn cứ','Operations':'M Vận hành dữ liệu và phát hành','Accounts':'N Tài khoản và quyền truy cập bổ sung','Workflow':'O Làm việc trong VS Code và kiến trúc API'}.get(a.get('group'))
    if not group:
        group=('L Pikafish và phân tích có căn cứ' if a['id'].startswith('N-ENG') else 'M Vận hành dữ liệu và phát hành' if a['id'].startswith('N-OPS') else 'N Tài khoản và quyền truy cập bổ sung' if a['id'].startswith('N-ACC') else 'O Làm việc trong VS Code và kiến trúc API')
    new.append(dict(title=a['title'],body=a['question'],fields=a['answer_fields'],proposal=a['proposal'],priority=a['priority'],related=[x.replace('V5-','V5 ') for x in a['related_v5']],source='Trao đổi mới nhất về engine và VS Code; '+a.get('why',''),orig_id=a['id'],group=group,condition=a.get('skip_if','')))
for i,q in enumerate(new,1): q['id']=f'BS {i:02}'
new[27]['fields'].append('Có ghi thống kê nước sai luật không và cách phân biệt bấm nhầm với chưa hiểu luật')
mapping={q['orig_id']:q['id'] for q in new}
mapping['N-ENG07']=mapping['PA13'];mapping['N-ENG09']=mapping['PA09']
assert len(old)==46 and len(new)==50
allq=old+new
groups=[]
for g in V['groups']:
    ids=[x.replace('V5-','V5 ') for x in g['ids']]
    groups.append((g['id']+' '+g['title'],[q for q in old if q['id'] in ids]))
for name in dict.fromkeys(q['group'] for q in new):groups.append((name,[q for q in new if q['group']==name]))

doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.65);s.bottom_margin=Inches(.62);s.left_margin=s.right_margin=Inches(.72)
s.header_distance=s.footer_distance=Inches(.28)
for name in ['Normal','List Bullet']:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.02;st.paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
hp=s.header.paragraphs[0];hp.text='CỜ CHILL   |   BỘ CÂU HỎI HỢP NHẤT V5   |   21/10/2026'
hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(0,0,0)
ft=s.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT
ft.add_run('Bản để điền   •   Trang ').font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');ft._p.append(f)
doc.core_properties.title='Cờ Chill bộ câu hỏi hợp nhất V5'
doc.core_properties.subject='Chốt yêu cầu phát triển từ bản web đã phát hành'
doc.core_properties.author='Cờ Chill'
bid=0
def para(t='',bold=False,size=None):
    p=doc.add_paragraph();r=p.add_run(t);r.bold=bold
    if size:r.font.size=Pt(size)
    return p
def heading(t,level=1):return doc.add_heading(re.sub(r'[^\w\sÀ-ỹ]','',t),level)
def page(t):doc.add_page_break();return heading(t)
def mark(p,name):
    global bid
    bid+=1;a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bid));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bid));p._p.insert(0,a);p._p.append(b)
def link(p,text,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    else:
        rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'174B70');pr.append(co)
    r.append(pr);tx=OxmlElement('w:t');tx.text=text;r.append(tx);h.append(r);p._p.append(h)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for ed in ['top','bottom','left','right']:
                e=OxmlElement('w:'+ed);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            cp.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233D53' if ri==0 else 'F2F5F8' if ri%2==0 else 'FFFFFF');cp.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(9.5);r.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    para('')
    return t

heading('Cờ Chill bộ câu hỏi hợp nhất V5',0)
para('Một file để chốt bản web tiếp theo',True)
para('96 câu: giữ đủ 46 câu V5 và thêm 50 câu sau khi gộp phần trùng. Các nhóm có số câu khác nhau theo nhu cầu. File này thay các biểu mẫu V5 trước để bạn trả lời; tài liệu gốc vẫn được giữ nguyên.')
para('Mục tiêu là làm rõ các quyết định còn thiếu để tiếp tục đúng hướng của bạn: học từ ván thật, suy nghĩ trước đáp án, phân tích có căn cứ, thực hành tiếp, lịch ôn, tài khoản nhiều thiết bị và quản trị. Phần mới bổ sung Pikafish, giao diện từng màn, vận hành và cách làm việc trong VS Code.')
heading('Cách điền',2)
for t in ['Điền ngay dưới mỗi câu. Dùng số 1, 2, 3 tương ứng với các mục Cần điền; thêm dòng/ảnh tùy ý. Không cần vừa chỗ trống hoặc tự viết FEN, code, SQL hay chọn framework.', 'Nếu đồng ý toàn bộ, ghi Giữ đề xuất. Nếu chỉ đồng ý một phần, ghi rõ phần giữ và phần đổi. Không có lựa chọn nào được điền sẵn thay bạn.', 'Nếu chưa biết, ghi Chưa rõ và mô tả kết quả mong muốn hoặc Giao kỹ thuật đề xuất rồi gửi tôi duyệt. Bỏ trống không đồng nghĩa đồng ý.', 'V5 01–46 giữ mã cũ. Nếu bạn đã có câu trả lời ở file khác, chép phần đó vào đúng mã rồi chỉ bổ sung các mục còn thiếu. Không phải điền hai file.', 'Câu Chỉ khi triển khai được ghi Chưa làm đợt này nếu không thuộc bản tới. Tài khoản, đồng bộ và admin đã thuộc V5; lựa chọn này không tự bỏ phần đã chốt.', 'Không ghi mật khẩu, token, API key, mã khôi phục hoặc dữ liệu riêng của người thử. Chỉ ghi tên dịch vụ, vai trò và file mẫu đã loại dữ liệu riêng.']:
    doc.add_paragraph(t,style='List Bullet')
para('Người trả lời: ____________________    Ngày bắt đầu: ______________\nMốc mong muốn có bản thử: __________    Ngày gửi lại: ______________')
para('Đề xuất trong file là để bạn quyết định. Ghi đồng ý thiết kế không tự cho phép mua dịch vụ, đổi public, push/deploy hoặc xóa dữ liệu; các hành động cần quyền vẫn xin duyệt trên bản cụ thể.',True)

page('Hiện trạng dùng làm điểm xuất phát')
para('Mốc phát hành có hồ sơ gần nhất: Site phiên bản 3 ngày 16/10/2026, source ứng dụng ở online-preview. URL đã phát hành:')
link(para(),'Mở địa chỉ Cờ Chill đã ghi trong báo cáo',url='https://co-chill-ngay-mai.cosmicanole32.chatgpt.site/')
para('Đây là mốc theo báo cáo phát hành và source local được đọc, không phải tuyên bố vừa kiểm thử lại toàn bộ website công khai ngày 21/10. Nếu bạn đã sửa một bản mới ở nơi khác, ghi tại '+mapping['N-WORK01']+'.')
table(['Đã có trong bản được ghi nhận','Chưa có hoặc còn giới hạn'],[
('Chơi máy cơ bản/hai người cùng máy, luật hợp lệ, ký phổ và nhập text.','Chưa có PvP online; luật phân xử lịch sử chưa đầy đủ cho mọi trường hợp giải đấu.'),
('Xem lại, phản hồi trước lời giải, gợi ý có kiểm soát, thử lại và bài liên quan.','Engine tự viết tối đa hai nửa nước; chưa có Pikafish hoặc LLM chấm lập luận.'),
('Nháp, giữ chỗ dở, lịch ôn và tiến độ trong trình duyệt.','Chưa có tài khoản, đồng bộ server và dashboard admin. Đây là yêu cầu V5 đã chốt, chưa phải tính năng đang chạy.'),
('12 bài tự dựng; ba theme, chữ quân Hán/Việt và giao diện responsive.','Chưa có kho bài lớn được chuyên gia duyệt, không có bằng chứng tác động học tập đã kiểm định.'),
('Xuất dữ liệu, xóa dữ liệu local theo các thao tác hiện có.','Xuất chưa tự phục hồi; chưa giữ hướng lật bàn qua reload; đồng hồ lúc kết thúc ván còn lỗi hiển thị đã ghi nhận.')
],[3.5,3.5])
heading('Các nguyên tắc tiếp tục giữ',2)
for t in V['locked_decisions_do_not_reask']:
    doc.add_paragraph(t,style='List Bullet')
para('Đối với chi phí, 100.000 đồng/tháng mong muốn và 200.000 đồng/tháng trần là câu trả lời trước đây. Nếu muốn thay đổi do engine mạnh, chỉ cập nhật tại V5 37. Không coi việc bạn mua công cụ AI lập trình là đã tăng ngân sách vận hành website.')

page('Thứ tự trả lời và mục lục')
mark(doc.paragraphs[-1],'Contents')
para('Mức 1: cần chốt trước khi làm phần liên quan. Mức 2: tinh chỉnh để nghiệm thu tốt hơn. Mức 3: chỉ cần chốt khi đưa tính năng đó vào phạm vi. Bạn có thể trả lời theo chủ đề, không bắt buộc theo số trang.')
para('Ưu tiên trước: V5 01–04, 08–10, 14–16, 25–26, 28–30, 32, 35–37, 39–46; '+mapping['N-ENG01']+'–'+mapping['N-ENG06']+'; '+mapping['N-ACC04']+' và '+mapping['N-WORK01']+'. Sau đó hoàn thiện toàn bộ các câu còn lại hoặc ghi rõ để sau.')
for gi,(name,qs) in enumerate(groups):
    pp=para();link(pp,f'{name}   |   {len(qs)} câu   |   {qs[0]["id"]} đến {qs[-1]["id"]}',anchor=f'Group{gi}')
para('Phụ lục cuối file: bản đồ từng màn hình, mẫu cung cấp thế cờ và lỗi giao diện, các lựa chọn đã chốt, nguồn đối chiếu và danh sách điều kiện phát hành. Đây là phần hỗ trợ điền, không phải biểu mẫu thứ hai.')
heading('Những việc kỹ thuật được giao cho người triển khai',2)
para('Chuyển FEN/tọa độ, quản lý tiến trình UCI, cấu trúc API/database, cache/hàng đợi, giới hạn tài nguyên, kiểm tra biến hợp lệ, xác thực/phân quyền, giữ phiên bản và test là việc của kỹ thuật. Bạn quyết định trải nghiệm, dữ liệu được phép dùng, giới hạn chi phí và kết quả cần đạt. Chưa có câu trả lời không cho phép kỹ thuật tự bịa nội dung cờ hoặc thay nguyên tắc đã chốt.')

question_locations=[]
for gi,(name,qs) in enumerate(groups):
    page(name);mark(doc.paragraphs[-1],f'Group{gi}')
    pp=para();pp.paragraph_format.space_after=Pt(7);link(pp,'Về mục lục',anchor='Contents')
    if name.startswith('L '):
        para('Pikafish được đề xuất tích hợp như bộ tính toán; không phải engine đã cài. Mã nguồn GPL-3.0 và điều kiện trọng số NNUE phải kiểm tra riêng. Nguồn chính thức của trọng số hiện ghi không dùng thương mại khi chưa có phép; cần đối chiếu đúng file được chọn. Nguồn ở cuối tài liệu.',size=10)
    if name.startswith('J '):para('Nhóm này làm rõ luồng và các trạng thái còn thiếu. Hình thức màu/logo chốt ở V5 20; phần K chốt mẫu thiết kế và màn tài khoản. Không cần trả lời lại dữ liệu hoặc quyền đã có ở nhóm I.',size=10)
    for idx,q in enumerate(qs):
        if idx and idx%2==0:doc.add_page_break()
        title=heading(q['id']+' '+q['title'],2);mark(title,q['id'].replace(' ','_'))
        question_locations.append(q['id'])
        rel=' · Liên quan: '+', '.join(q['related']) if q.get('related') else ''
        mp=para(f'Mức {q["priority"]}'+rel,size=9);mp.paragraph_format.keep_with_next=True
        if q.get('condition') and (q['priority']==3 or q.get('orig_id','').startswith('N-')):para('Điều kiện áp dụng: '+q['condition'],size=9)
        body=para(q['body']);body.paragraph_format.keep_with_next=True
        pp=para('Đề xuất chưa duyệt: '+q['proposal']);pp.paragraph_format.keep_with_next=True
        fields='; '.join(f'{i+1}. {v}' for i,v in enumerate(q['fields']))
        fp=para('Cần điền: '+fields,size=10);fp.paragraph_format.keep_with_next=True
        ap=para('Trả lời của bạn: [Điền theo số ở trên hoặc ghi Giữ đề xuất / Chưa rõ / Để sau kèm lý do]')
        ap.paragraph_format.space_after=Pt(16)
        # The editable response paragraph expands in Word as the owner types.

page('Bản đồ các màn hình cần đối chiếu')
para('Bảng này giúp không bỏ quên trang. Không phải yêu cầu bạn thiết kế lại từ đầu. Nếu giữ giao diện đã chọn, ghi Giữ hiện tại tại câu liên quan; nếu sửa, dùng ảnh/khung phác và mã màn hình. Các route là tên trong source đang đọc, không chứng minh có tính năng tài khoản online.')
rows=[]
for a in P['route_coverage']:
    refs=[f'V5 {int(x):02}' for x in a['existing_v5']]+[mapping[x] for x in a['new_questions']]
    rows.append((a['label']+'\n'+a['route'],', '.join(refs)))
table(['Màn hình hoặc trạng thái','Câu chốt yêu cầu'],rows,[3.4,3.6])
para('Mỗi màn cần kiểm tra: trạng thái bình thường, chưa có dữ liệu, đang tải, lỗi, mất mạng, thiếu quyền, thao tác thành công và dữ liệu chưa lưu. Mặc định: nói đúng tình trạng, có bước tiếp, giữ nháp khi phù hợp, không giả dữ liệu hoặc phần trăm tiến độ. Ngoại lệ điền tại '+mapping['PA25']+'.')

page('Mẫu gửi ví dụ và tiêu chí nghiệm thu')
heading('Một hình cờ hoặc ván mẫu',2)
para('Tên file nên bắt đầu bằng mã câu, ví dụ V5_14_hinh_01.png. Không cần biết FEN; ảnh rõ 90 giao điểm và bên đến lượt là đủ để kỹ thuật dựng lại rồi xác minh với bạn. Với luật lặp/chiếu/đuổi cần toàn bộ chuỗi nước liên quan, ảnh cuối đơn lẻ không đủ.')
para('Mã hình/van: __________   Thuộc câu: __________\nNguồn và quyền dùng: __________________________________________\nBên đi, mục tiêu bài, đoạn lịch sử cần xét: _________________________\nNước chấp nhận và vì sao: ______________________________________\nNước không đạt và vì sao: ______________________________________\nLời giải của bạn hoặc ghi Cần kiểm chứng: _________________________\nTên/chủ đề kỹ năng, độ khó dự kiến: ______________________________')
heading('Một vấn đề giao diện',2)
para('Mã câu / trang / thiết bị / trình duyệt / zoom: ______________________\nBước thao tác: ________________________________________________\nThực tế đang xảy ra: ___________________________________________\nBạn muốn thay đổi thành: _______________________________________\nẢnh hoặc bản phác: ____________________________________________\nCách nhận biết đã sửa xong: _____________________________________')
heading('Các điều kiện nền phải kiểm tra khi phát hành',2)
for t in ['Không cho nước sai luật; mỗi biến engine được phát lại qua luật. Nếu các bộ luật bất đồng, dừng và báo cần kiểm chứng.', 'Không lộ đáp án trước phản hồi hoặc coaching trong ván; giữ câu đầu và phân biệt có hỗ trợ với tự giải.', 'Không mất nháp hoặc tính hai lần khi gửi lại; lịch ôn bắt đầu từ hoàn thành hình và không được thay đổi bởi tự đánh giá.', 'Tài khoản tạo trên A thấy đúng trong admin B; học trên A thấy tiến độ đã đồng bộ khi đăng nhập C.', 'A không truy cập dữ liệu B; người thường không gọi chức năng admin; khóa và quyền xem có hiệu lực phía server.', 'Engine lỗi có trạng thái rõ và phục hồi; hết hạn mức không làm biến mất dữ liệu đã có. Nhãn engine, phiên bản và dữ liệu thật phải đúng.', 'Xuất/xóa và phục hồi được thử; dữ liệu test tách khỏi pilot; báo cáo không bịa người dùng hoặc hiệu quả.', 'Bạn xem bản thử có changelog, test, giới hạn và cách quay lại trước khi duyệt push/deploy.']:
    doc.add_paragraph(t,style='List Bullet')

page('Các lựa chọn đã có để không phải trả lời lại')
para('Đây là sổ tóm tắt ý định đã nêu, không phải danh sách tính năng hoàn tất. Các chi tiết còn mâu thuẫn được hỏi ở phần chính. Nếu muốn đổi một lựa chọn, ghi mã C tương ứng và quyết định mới; không cần chép lại cả lịch sử.')
for c in P['confirmed_owner_choices']:
    ch=para(c['id']+' '+c['topic'],True)
    ch.paragraph_format.keep_with_next=True
    ch.paragraph_format.space_before=Pt(5)
    ch.paragraph_format.space_after=Pt(2)
    cb=para(c['decision'])
    cb.paragraph_format.keep_with_next=True
    cb.paragraph_format.space_after=Pt(2)
    cs=para('Nguồn: '+', '.join(c['source_ids']),size=9)
    cs.paragraph_format.space_after=Pt(3)

page('Nguồn đối chiếu và cách gửi lại')
para('Repository gốc: E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai. Các đường dẫn dưới đây dùng R để chỉ đúng thư mục này; chữ R không phải biến terminal.')
for label,path in [
('V5 cập nhật đủ 46 câu','R\\docs\\CO_CHILL_V5_QUESTIONS.md'),
('Bản Word V5 tài khoản','R\\deliverables\\Co_Chill_V5_Cap_nhat_Tai_khoan.docx'),
('V4 đã trả lời',r'E:\ThanhTQ\App cờ tướng\10-14\V4.docx'),
('UI UX và Luật đã trả lời',r'E:\ThanhTQ\App cờ tướng\10-13\UI UX + Luat.docx'),
('Sổ quyết định và phạm vi luật','R\\docs\\UIUX_DECISION_REGISTER_2026_10_15.md\nR\\docs\\RULES_UIUX_2026_10_15.md'),
('Báo cáo bản phát hành','R\\BAO_CAO_UIUX_V5_2026_10_16.md'),
('Đặc tả nền','R\\docs\\spec\\01_PRODUCT_SPEC.md\nR\\docs\\spec\\07_LEARNING_MODEL_SPACED_REPETITION.md'),
('Cách tự review','R\\deliverables\\Co_Chill_Huong_dan_review_va_update_2026_10_20.docx')]:
    para(label,True);para(path,size=10)
para('Các trao đổi ngày 21/10 về VS Code, engine mạnh và yêu cầu gom một file được đưa vào nhóm bổ sung. Phần chữ V4 và UI UX được đọc đầy đủ, kể cả bảng; ảnh thiết kế được đối chiếu qua hồ sơ đã có, không có một vòng đánh giá lại toàn bộ ảnh gốc trong lần tổng hợp này.')
heading('Nguồn kỹ thuật tham khảo cho Pikafish',2)
for text,url in [('Mã nguồn và giấy phép engine','https://github.com/official-pikafish/Pikafish'),('Giao thức UCI','https://github.com/official-pikafish/Pikafish/wiki/UCI-%26-Commands'),('Điều kiện trọng số NNUE','https://github.com/official-pikafish/Networks')]:link(para(),text,url=url)
para('Các nguồn engine được tra ngày 21/10/2026. Điều kiện áp dụng cần kiểm tra lại với đúng phiên bản khi triển khai. UCI, độ sâu và tham số kỹ thuật không phải câu hỏi bắt buộc bạn tự quyết.')
heading('Chốt một lần trước khi bắt đầu code',2)
para('Bạn gửi lại chính file này đã điền và file mẫu đặt tên theo mã câu. Không cần gửi lại V3/V4 nếu không có thay đổi. Sau khi nhận, người triển khai phải tổng hợp: quyết định cuối cùng, mâu thuẫn còn lại, phạm vi đợt đầu, tiêu chí nghiệm thu, chi phí dự kiến và hành động cần duyệt. Câu để trống hoặc lời đề xuất chưa được đồng ý không được biến thành quyết định của bạn.')
para('Xác nhận cuối: tôi đã điền các phần cần làm trong đợt tới và ghi rõ phần chưa chốt.\nNgười trả lời: ____________________   Ngày: ____________________\nCác câu vẫn cần trao đổi: _______________________________________')

for part in [doc.styles.element,doc.element]:
    for node in part.xpath('.//w:pBdr'):node.getparent().remove(node)
doc.save(OUT)
report={'output':str(OUT),'questions':len(allq),'v5_count':len(old),'new_count':len(new),'groups':len(groups),'question_ids':question_locations,'mapping':mapping,'answer_field_count':sum(len(q['fields']) for q in allq),'sources_unchanged':True,'old_questions_preserved':len(old)==46,'authoring_only':True}
(W/'question_bank.json').write_text(json.dumps({'groups':groups,'questions':allq},ensure_ascii=False,indent=2),encoding='utf-8')
(W/'build_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['mapping','question_ids']},ensure_ascii=False,indent=2))
