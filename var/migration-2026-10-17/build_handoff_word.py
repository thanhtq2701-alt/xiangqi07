from pathlib import Path
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT/'deliverables/Co_Chill_So_tay_Code_VSCode_2026_10_17.docx'
PKG = ROOT/'deliverables/Co_Chill_VSCode_2026_10_17'
doc = Document()
sec = doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(1.9); sec.right_margin=Cm(1.9); sec.footer_distance=Cm(.75)
for name,size,bold in [('Normal',10.5,False),('Title',25,True),('Subtitle',12,False),('Heading 1',17,True),('Heading 2',12,True),('Heading 3',11,True),('List Bullet',10.5,False)]:
    s=doc.styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(6); s.paragraph_format.line_spacing=1.1
    for c in s.element.xpath('.//w:color'):
        for attr in list(c.attrib):
            if 'theme' in attr.lower():del c.attrib[attr]
    if name.startswith('Heading'):s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(9)
for b in doc.styles.element.xpath('.//w:pBdr'):b.getparent().remove(b)
doc.core_properties.title='Sổ tay mã nguồn Cờ Chill và chuyển sang VS Code'
doc.core_properties.author='Cờ Chill'
doc.core_properties.subject='Bàn giao mã nguồn và kiến trúc cho chủ sản phẩm'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Cờ Chill  |  17/10/2026  |  ').font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)

def p(t,style=None):
    v=doc.add_paragraph(t,style);v.paragraph_format.widow_control=True;return v
def h(t):return p(t,'Heading 2')
def page(n,t):
    v=p(f'{n}  {t}','Heading 1');v.paragraph_format.page_break_before=True;return v
def code(t):
    v=p(t);v.paragraph_format.space_before=Pt(3);v.paragraph_format.space_after=Pt(8)
    for r in v.runs:r.font.name='Consolas';r.font.size=Pt(9)
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F4F5');v._p.get_or_add_pPr().append(shade)
    return v
def bullet(t):return p(t,'List Bullet')
def table(headers,rows,widths):
    tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
    for i,w in enumerate(widths):tb.columns[i].width=Cm(w)
    for i,s in enumerate(headers):tb.rows[0].cells[i].text=s
    for row in rows:
        cells=tb.add_row().cells
        for i,s in enumerate(row):cells[i].text=s
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    tb._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(tb.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'24483F' if ri==0 else ('F3F6F5' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1.05
                for rr in pp.runs:rr.font.name='Arial';rr.font.size=Pt(9.5);rr.font.bold=ri==0;rr.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    p('')
    return tb
def filecard(title,name,role,missing,edit):
    h(title);code('online-preview/'+name)
    p('Vai trò: '+role);p('Nếu thiếu: '+missing);p('Khi cần sửa: '+edit)

p('Sổ tay mã nguồn Cờ Chill\nvà chuyển sang VS Code','Title')
p('Bản bàn giao ngày 17 tháng 10 năm 2026','Subtitle')
p('Bạn có thể mở và phát triển Cờ Chill trong VS Code với trợ lý lập trình bạn chọn. Website hiện tại dùng HTML, CSS và JavaScript thuần; việc chạy web không phụ thuộc vào Codex hoặc Claude. Bộ nguồn gọn gồm 29 file, trong đó có 10 tài nguyên chạy web; phần còn lại phục vụ kiểm thử và bảo trì.')
h('Bạn nên dùng bộ nào')
p('Dùng thư mục dưới đây làm nơi làm việc mới. Sao chép cả thư mục nếu chuyển máy; không ghép code vào một file và không chỉ lấy index.html.')
code(str(PKG))
p('Bản ZIP để chuyển máy nằm cùng thư mục deliverables, tên Co_Chill_VSCode_2026_10_17.zip. Giải nén trước khi mở bằng VS Code. Word này là tài liệu đọc; code thực thi nằm trong thư mục hoặc ZIP, không nằm trong Word.')
h('Ba việc làm ngay')
bullet('Mở thư mục mới trong VS Code bằng File > Open Folder. Đọc README.md, rồi xem online-preview/index.html, style.css và app.js.')
bullet('Mở Terminal > New Terminal và chạy lệnh ở mục 2. Kiểm tra trang chủ, một nước cờ và một bài học trước khi sửa.')
bullet('Gửi AI_HANDOFF.md cùng yêu cầu cụ thể cho trợ lý mới. Với nâng cấp V5, cung cấp thêm câu trả lời V5 của bạn; tài khoản thật vẫn là phần cần triển khai.')
h('Mốc nguồn đã đối chiếu')
p('10 tài nguyên chạy web trong online-preview/ khớp từng byte với var/site-preview/dist/ của checkout phát hành sạch tại commit 55bd70cd4c19c52192d5d80160d13982e1c7a80a. Dịch vụ Sites báo phiên bản mới nhất là 3. Repo đầy đủ đang ở commit 0c9f20b2a81033193cd7bf81bcd4aeacc7d31c5a và có thay đổi tài liệu chưa commit.')
p('Nguồn cục bộ đã được đối chiếu với checkout phát hành. Yêu cầu tải trực tiếp các tài nguyên công khai qua HTTP trong lượt kiểm tra này bị trả 403, nên không coi đó là phép xác minh byte của tài nguyên đang được hosting phục vụ.')
p('Cách đọc: mục 2 giúp chạy ngay; mục 3–8 giải thích kiến trúc; mục 9–12 giải thích dữ liệu và file có thể tách; mục 13–16 hướng dẫn sửa, kiểm thử, Git và bàn giao.')

page(2,'Mở trong VS Code và chạy trên máy khác')
p('Trong VS Code, mở đúng thư mục chứa README.md và thư mục con online-preview. Một cửa sổ editor trống không tự có mã nguồn, server hay tài khoản hosting. Đây là cách làm việc theo thư mục được VS Code hỗ trợ [1].')
code('python -m http.server 8769 --bind 127.0.0.1 --directory online-preview')
p('Sau khi terminal báo đang phục vụ, mở http://127.0.0.1:8769/#home. Giữ terminal chạy; Ctrl+C để dừng. Nếu dùng Python Launcher trên Windows, thay python bằng py. Python ở đây chỉ gửi file tĩnh; không chạy server.py của dự án [2].')
p('Nếu mở trực tiếp online-preview làm thư mục gốc, bỏ phần --directory online-preview khỏi lệnh. Với bộ bàn giao, nên mở thư mục cha để thấy cả kiểm thử và README.')
h('Máy hiện tại đã có gì')
table(['Thành phần','Kết quả kiểm tra và việc cần làm'],[
('Python','PATH hiện dùng Python 3.14.0; đủ chạy web. Bộ nguồn không yêu cầu pip install.'),
('VS Code','Đã tìm thấy ở E:\\ThanhTQ\\Code\\Microsoft VS Code. Có thể mở thư mục mới độc lập với repo cũ.'),
('Node.js','Lệnh node hiện trỏ đến bản 6.14.0 ở E:\\ThanhTQ\\Code\\node.exe. Quá cũ cho kiểm thử này. Dùng Node 22 trở lên hoặc truyền đường dẫn Node hiện đại vào checker.'),
('Trình duyệt','Chạy JavaScript, vẽ bàn cờ và lưu tiến độ. Node không cần thiết để người chơi dùng bản web này.'),
('Claude hoặc trợ lý khác','Hỗ trợ đọc và sửa mã. Không phải thư viện được website gọi; đổi trợ lý không đòi viết lại website.')
],[3.1,14.1])
h('Lỗi thường gặp')
bullet('Không nhận lệnh python: thử py --version; nếu chưa có Python, cài từ nguồn chính thức rồi mở terminal mới.')
bullet('Cổng 8769 đang bận: dùng server đã có nếu đúng dự án; hoặc đổi sang 8770. Đổi cổng tạo vùng dữ liệu trình duyệt khác. Không tự tắt tiến trình lạ.')
bullet('Trang trống hoặc báo 404: kiểm tra thư mục terminal; phải có online-preview/index.html. Mở Developer Tools > Console để đọc lỗi JavaScript.')
bullet('Không dùng cách nhấp đúp HTML làm cách chạy chuẩn. Chạy qua HTTP giúp đường dẫn và vùng lưu dữ liệu nhất quán.')
p('Không cần npm install, package.json, API key, database hoặc dịch vụ trả phí để chạy bản hiện tại. Cần Node hiện đại chỉ khi chạy bộ kiểm thử JavaScript.')

page(3,'Kiến trúc tổng thể của website hiện tại')
p('Bản đang đưa lên ChatGPT Site là ứng dụng tĩnh có tương tác. Hosting gửi file; phần chơi cờ, phân tích nông và học tập chạy trên trình duyệt của từng người. “Tĩnh” ở đây nói về cách phục vụ file, không có nghĩa trang chỉ chứa hình ảnh và chữ.')
code('Hosting hoặc HTTP server local\n  |\n  +-- index.html               khung và thứ tự tải\n  +-- style.css + logo/icon    hình thức\n  +-- app.js                   màn hình và thao tác\n        |\n        +-- xiangqi.js         luật và tìm kiếm cờ\n        +-- game-tools.js      ký phổ và nhập ván\n        +-- fixtures.js        12 bài tập tự dựng\n        +-- learning.js        phiên học và lịch ôn\n        +-- experience.js      tùy chọn và thống kê\n        |\n        +-- localStorage      dữ liệu trên trình duyệt')
h('Các khái niệm cần phân biệt')
table(['Khái niệm','Trong dự án này'],[
('Frontend','Phần trình duyệt hiển thị và chạy: toàn bộ 10 tài nguyên trong online-preview/.'),
('Backend','Chương trình xử lý ở máy chủ. Bản Site hiện tại chưa có backend nghiệp vụ; server.py là một app local khác.'),
('API','Giao diện gọi chức năng. XQ và CC là API nội bộ của code. Bản Site không gọi API HTTP tài khoản hoặc phân tích từ server.'),
('Database','Kho dữ liệu tập trung chưa có trong Site. localStorage thuộc từng trình duyệt. SQLite của app Python là kho riêng.'),
('Deploy','Đưa một phiên bản file hoặc chương trình lên hosting. Save, commit và push không tự đồng nghĩa với deploy.')
],[3.1,14.1])
h('Thứ tự tải không được đảo tùy ý')
code('xiangqi → game-tools → fixtures → learning → experience → app')
p('index.html tải sáu script với defer. xiangqi tạo XQ; các module trợ giúp nối thêm chức năng vào CC; fixtures tạo FIXTURES. app.js dùng các thành phần đó sau cùng. Bản này chưa dùng import kiểu ES module, React, Next.js, TypeScript hoặc trình đóng gói.')

page(4,'Ba file điều khiển phần nhìn và thao tác')
p('Vị trí gốc của mỗi file dưới đây là E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai\\online-preview. Trong bản gọn, các file nằm trong online-preview/ bên trong thư mục mới ở trang 1. Tên và cấu trúc được giữ nguyên.')
filecard('Khung trang','index.html','Tạo header, sidebar, menu điện thoại, vùng main, thông báo và các thẻ tải CSS/script. Đây là cửa vào của toàn bộ website.','Không còn trang vào chuẩn hoặc các điểm gắn giao diện; JavaScript không tìm được phần tử cần dùng.','Menu cố định, tiêu đề tab ban đầu, metadata, logo và thứ tự tải tài nguyên. Không cần thêm một HTML cho mỗi trang nội bộ.')
filecard('Giao diện và bố cục','style.css','Định nghĩa màu, theme, chữ, khoảng cách, thẻ nội dung, form, bàn cờ, quân cờ và cách co giãn trên điện thoại.','Code có thể vẫn tải nhưng bố cục hỏng; bàn cờ và vị trí quân có thể không còn hiển thị dùng được.','Màu sắc, kích thước, khoảng cách, responsive và các trạng thái trực quan. Kiểm tra cả desktop và điện thoại sau sửa.')
filecard('Bộ điều phối màn hình','app.js','Đọc địa chỉ #, dựng nội dung vào main, xử lý nút bấm và form; nối luật, bot, review, bài học, đồng hồ, âm thanh và localStorage.','Chỉ còn khung tĩnh; phần chính không hiện hoặc không tương tác được.','Nội dung màn hình, quy trình bấm, hành vi chơi/học. Tìm render(), home(), playSetup(), learn(), settings() để bắt đầu đọc. Một số hàm đang viết gọn trên dòng dài; định dạng lại cần đi kèm kiểm thử.')

page(5,'Ba file về cờ và nội dung bài tập')
filecard('Luật và máy chơi','xiangqi.js','Biểu diễn bàn 90 giao điểm; đọc FEN; tạo nước hợp lệ cho bảy loại quân; kiểm tra an toàn Tướng, chiếu hết và hết nước; đánh giá vật chất và tìm kiếm nông tối đa hai nửa nước.','Chơi, kiểm tra luật và phân tích không hoạt động đúng. Không thay bằng một đoạn AI tự đoán nước đi.','Luật quân, kiểm tra FEN, giá trị đánh giá hoặc tìm kiếm. Chạy đối chiếu với luật Python và test các quân trước khi dùng kết quả mới.')
filecard('Ký phổ và công cụ ván','game-tools.js','Thêm chức năng vào CC: ký phổ Việt, nhập văn bản ván, lịch sử theo lượt, lý do nước cấm và xử lý lặp trong phạm vi đã kiểm chứng.','Ký phổ, nhập ván, xem lại và một số đường xử lý kết thúc ván lỗi; phần khác có thể hiện trước khi gặp lỗi.','Cú pháp nhập, hiển thị biên bản, giải thích nước cấm, case lặp. Đây chưa phải toàn bộ luật trọng tài giải đấu.')
filecard('Bộ bài đang được hiển thị','fixtures.js','Chứa 12 bài tổng hợp thuộc bốn nhóm kỹ năng, gồm FEN, nước mẫu, tập nước hợp lệ, câu giải thích và nguồn synthetic-codex-v1. App dùng trực tiếp biến FIXTURES.','Danh sách bài, bài liên quan và các màn hình tham chiếu bài mẫu sẽ hỏng. Đây là dữ liệu ứng dụng thật sự cần, dù tên gợi ý “fixture”.','Sửa nguồn bài ở app/puzzles.py và logic xuất ở tools/export_preview_fixtures.py, rồi tạo lại và kiểm tra. Nước mẫu không được quảng cáo là nước tốt duy nhất nếu chưa chứng minh.')
p('Không có engine Pikafish, trọng số mạng nơron hoặc bộ ván thương mại được đóng gói. Điểm máy là ước lượng tìm kiếm nông; không dùng để khẳng định trình độ thi đấu hay chất lượng chiến lược toàn diện.')

page(6,'Phiên học trải nghiệm và tài sản hình ảnh')
filecard('Trạng thái học và lịch ôn','learning.js','Tạo trạng thái ban đầu, lọc dữ liệu đã lưu, quản lý các bước reflect/explain/retry/complete, ghi hoàn thành, đề xuất bài và hẹn ôn từ lúc hoàn thành. Giữ bài dở và bản tự đánh giá.','Khởi tạo, tiếp tục bài, ghi tiến độ và lịch ôn bị lỗi. app.js gọi CC.fresh() ngay đầu.','Cách chuyển bước học, điều kiện hoàn thành, tránh ghi trùng, bài liên quan và lịch ôn. Lịch hiện là quy tắc gợi ý, không phải mô hình ML.')
filecard('Tùy chọn và bằng chứng học','experience.js','Kiểm tra tùy chọn giao diện; câu hỏi và hỗ trợ khi phản hồi; thống kê lần tự làm so với có hỗ trợ; xuất dữ liệu; cập nhật đồng hồ.','App có thể lỗi ngay ở CC.preferences(); các chức năng phản hồi, đồng hồ, thống kê và xuất cũng thiếu.','Theme, cỡ chữ, giọng mở đầu, ngưỡng thống kê, gợi ý và đồng hồ. Thay giao diện cần giữ nguyên điều kiện phản hồi trước đáp án.')
table(['Tài sản','Vị trí và tác dụng','Nếu bỏ'],[
('Logo','online-preview/logo-co-chill.png\nẢnh nhận diện Cờ Chill do chủ sản phẩm cung cấp.','Mất logo hoặc hiện ảnh hỏng. Logic cờ vẫn chạy.'),
('Icon tab','online-preview/favicon.svg\nBiểu tượng nhỏ trên tab trình duyệt.','Mất icon. Logic cờ vẫn chạy.')
],[2.5,9.4,5.3])
p('Giữ cả 10 tài nguyên để bản sao giống bản đang phát hành. Tám file đầu liên quan trực tiếp đến hoạt động hoặc hình thức dùng được; hai hình ảnh không quyết định luật nhưng cũng không nên bỏ để tiết kiệm. Âm thanh được tạo bằng Web Audio, không cần thư mục nhạc riêng.')

page(7,'Mỗi trang trong website nằm ở đâu')
p('Các địa chỉ bắt đầu bằng # là trang nội bộ của cùng index.html. app.js đọc phần sau #, chọn hàm tương ứng rồi thay nội dung main; style.css tạo hình thức cho tất cả. Vì vậy không tìm thấy play.html hoặc profile.html là bình thường.')
table(['Địa chỉ hoặc trang','Hàm trong online-preview/app.js và nội dung'],[
('#home','home() — trang chủ và lối vào các hoạt động.'),
('#play','playSetup() — chế độ chơi, màu quân, thời gian, nhập ván.'),
('#game/<id>','gameView() — bàn đang chơi, đồng hồ, biên bản và thao tác ván.'),
('#lessons','lessons() — danh sách bài, bộ lọc và nội dung học.'),
('#learn','learn() — suy nghĩ, hướng dẫn, thử lại và tiếp tục bài.'),
('#opening/<số>','opening() — nội dung khai cuộc mẫu trong phạm vi app.'),
('#mistakes','mistakes() — danh sách ván của tôi và đường mở xem lại.'),
('#moments/<id>','moments() — các vị trí đáng xem lại của một ván.'),
('#replay/<id>','replay() — phát lại ván theo nước.'),
('#review','reviews() — bài đến lịch ôn.'),
('#summary','summary() — tự đánh giá và kết thúc phiên.'),
('#profile','profile() — hồ sơ và tiến độ tại trình duyệt này.'),
('#settings','settings() — tùy chọn, xuất và xóa dữ liệu cục bộ.'),
('#howto và #rules','howto(), rulesPage() — hướng dẫn thao tác và luật.'),
('#path','learningPath() — hành trình học gợi ý.'),
('#about','about() — mô tả bản thử và giới hạn.')
],[4.5,12.7])
p('Khi đổi nội dung của một trang, tìm hàm ở cột phải trước. Khi đổi menu chung, đọc index.html. Khi đổi màu hoặc khoảng cách, đọc style.css. Màn hình hồ sơ hiện chưa phải trang quản trị tài khoản.')

page(8,'Vòng học hiện hoạt động qua những phần nào')
table(['Bước','Code chịu trách nhiệm và điều cần giữ'],[
('1  Chơi hoặc nhập','app.js điều phối; xiangqi.js kiểm tra luật; game-tools.js đọc ký phổ. Có máy local và hai người cùng thiết bị.'),
('2  Xem lại sau ván','app.js chặn coaching khi ván còn diễn ra. Phân tích tối đa 16 lượt gần đây phù hợp với bên người học và chọn tối đa ba vị trí.'),
('3  Tự suy nghĩ','learning.js và experience.js yêu cầu ba phản hồi cùng nước dự định, hoặc luồng xin trợ giúp rõ ràng. Nháp được lưu.'),
('4  Chốt hướng dẫn','app.js trình bày WHAT, WHY, HOW từ dữ kiện hợp lệ và lời mẫu. Không có LLM tự sinh biến cờ hoặc tự chấm lập luận.'),
('5  Thử lại','Kiểm tra nước bằng luật, ghi có hỗ trợ/đã xem đáp án; không tính retry sau lời giải như thành công độc lập lần đầu.'),
('6  Bài liên quan','learning.js chọn cùng nhóm kỹ năng và ưu tiên ID chưa gặp. Bài bị bỏ qua có trạng thái chờ để tiếp tục.'),
('7  Ghi nhận và ôn','Hoàn thành từng hình tạo lịch riêng từ thời điểm kết thúc; tránh ghi nhận lặp. Chỉ mở bài chưa có nghĩa đã hoàn thành.'),
('8  Tự đánh giá','Bản tổng kết phiên tách khỏi bằng chứng đi quân. Không tự tăng mức thành thạo chỉ vì người dùng tích “đã hiểu”.'),
('9  Tiến độ','app.js và experience.js tổng hợp dữ liệu local; lưu bằng localStorage của trình duyệt đang dùng.')
],[3.5,13.7])
h('Những điều không nên hứa với người dùng')
bullet('Engine nông có thể bỏ lỡ phương án dài. Không phải mọi ván đều có sai lầm mà hệ thống tìm được.')
bullet('12 bài là nội dung tự dựng, chưa phải giáo trình lớn hoặc bộ dữ liệu đã chứng minh hiệu quả học.')
bullet('Ưu tiên chưa gặp là theo ID trong dữ liệu trình duyệt; chưa có lịch sử chung từ mọi thiết bị.')
bullet('Đáp án nằm trong JavaScript công khai. Khóa trước suy nghĩ giúp quá trình học, không bảo mật đề thi trước người đọc source.')

page(9,'Dữ liệu nào không đi theo khi copy code')
p('Mã nguồn là công thức của ứng dụng. Tiến độ là dữ liệu sinh ra khi chơi và học. Bộ source không chứa ván, câu trả lời hay dữ liệu riêng của người chơi.')
table(['Nơi lưu','Nội dung và ảnh hưởng'],[
('co-chill-v4-browser','Khóa localStorage: ván, phiên học, câu trả lời, bản ghi bài, lịch ôn, tổng kết và tên hiển thị.'),
('co-chill-preferences-v1','Khóa localStorage: theme, chữ quân, cỡ chữ, âm thanh, sidebar, trình độ tự khai và giọng.'),
('Bản Python riêng','var/chill.sqlite3 lưu dữ liệu phiên khách của server.py. Không phải database của Site và không đồng bộ với Site.')
],[5.5,11.7])
h('Vì sao mở chỗ khác thấy như tài khoản mới')
p('localStorage tách theo trình duyệt, hồ sơ và origin gồm giao thức, tên miền, cổng. http://127.0.0.1:8769, http://localhost:8769, http://127.0.0.1:8770 và tên miền chatgpt.site không dùng chung dữ liệu. Tên hiển thị hoặc avatar hiện là hồ sơ local, chưa phải tài khoản đăng ký.')
p('Nếu giữ cùng trình duyệt và cùng origin trên máy cũ, việc đổi thư mục source phục vụ có thể vẫn đọc dữ liệu cũ. Hãy xuất bản sao trước khi thử code có thay đổi cách lưu. Đóng hoặc ngừng dùng bản cũ khi chuyển để tránh hai bản cùng ghi lên một origin.')
h('Sao lưu hiện hỗ trợ đến đâu')
p('Vào Cài đặt để xuất JSON trước khi xóa dữ liệu hoặc chuyển môi trường. File xuất có thể đọc và lưu riêng; app chưa có chức năng nhập lại bản xuất cá nhân để phục hồi tự động. Không đưa file xuất này, browser profile hoặc SQLite vào repo công khai.')
h('Nâng cấp tài khoản trong V5')
p('Yêu cầu đã chốt gồm đăng nhập thật, lưu tập trung, đồng bộ nhiều thiết bị, thống kê người đăng ký và trang quản trị có phân quyền. Bộ nguồn hiện tại chưa triển khai các phần đó. Chúng cần backend, cơ sở dữ liệu chung và kiểm tra quyền phía server; một trang HTML quản trị riêng không đủ.')
p('Số người đăng ký phải đến từ bản ghi tài khoản thật. Không dùng tên local, số lượt truy cập hoặc số lần mở trang thay cho người dùng duy nhất. Việc chuyển tiến độ local vào tài khoản cần lựa chọn và xác nhận, nhất là máy dùng chung.')

page(10,'Những file giữ lại để sửa code an toàn')
p('Các file trang này không được trình duyệt tải. Bỏ chúng không làm bản JavaScript đã xuất ngừng chạy ngay, nhưng mất khả năng kiểm tra hoặc tạo lại nội dung. Bộ gọn giữ chúng để trợ lý mới tiếp tục phát triển mà không bỏ hàng rào kiểm thử.')
table(['Đường dẫn trong bộ gọn','Vai trò và hậu quả nếu bỏ'],[
('app/__init__.py','Đánh dấu package Python. Giữ cấu trúc nhập module ổn định cho công cụ tham chiếu.'),
('app/rules.py','Luật Python dùng để tạo bài và so sánh với JavaScript. Bỏ sẽ làm exporter, checker, parity và test luật lỗi.'),
('app/engine.py','Engine tham chiếu cho phép so kết quả tìm kiếm. Bỏ làm phần đối chiếu engine không chạy.'),
('app/puzzles.py','Nguồn 12 bài tự dựng. fixtures.js vẫn chạy nếu đã có, nhưng không còn nguồn để tạo lại và đối chiếu.'),
('tests/__init__.py','Đánh dấu package kiểm thử. Không liên quan runtime browser; giữ tổ chức test.'),
('tests/test_rules.py','Kiểm thử luật quân, FEN và trạng thái ván. Bỏ thì mất bộ phát hiện hồi quy ở luật tham chiếu.'),
('tools/export_preview_fixtures.py','Xuất fixtures.js từ nguồn bài và luật; chỉ chạy khi chủ động thay nguồn bài.'),
('tools/check_online_preview.py','Kiểm tra 12 bài, nguồn gốc, tập nước hợp lệ và nội dung xuất có khớp nguồn không.'),
('tools/check_browser_rules.py','So luật và engine JS với Python trên các thế kiểm tra và ván sinh xác định; cần Node hiện đại.'),
('tools/test_preview_learning.cjs','Kiểm tra reflection gate, ghi hoàn thành, lịch ôn, đề xuất, resume và dữ liệu lưu.'),
('tools/test_preview_game_tools.cjs','Kiểm tra ký phổ, nhập ván, lý do nước cấm và xử lý lặp.'),
('tools/test_preview_experience.cjs','Kiểm tra tùy chọn, trợ giúp, đồng hồ, export, thống kê và một số hành vi app.'),
('tools/check_all.py','Công cụ mới gom kiểm tra cú pháp và các suite ở trên; tổng hợp lỗi và trả mã thoát khác 0 nếu có lỗi.')
],[7.1,10.1])
p('app/ trong bộ gọn chỉ có bốn file tham chiếu. Nó không phải bản backend đầy đủ: không có service.py, storage.py, learning.py hoặc server.py. Không nhầm sự có mặt của app/ với việc tài khoản/SQLite đã hoạt động.')

page(11,'File hỗ trợ VS Code và cấu trúc bộ bàn giao')
table(['File','Tác dụng','Nếu bỏ'],[
('README.md','Hướng dẫn chạy, test, phạm vi và dữ liệu. Đây là nơi đọc đầu tiên.','Web vẫn chạy; người tiếp nhận dễ nhầm cách dùng.'),
('.gitignore','Loại cache, dữ liệu, env và file sinh khỏi bước Git add thông thường.','Web vẫn chạy; dễ commit nhầm file. Không tự xóa secret đã commit.'),
('.vscode/tasks.json','Tác vụ chạy web và kiểm thử qua menu Terminal > Run Task.','Vẫn dùng lệnh terminal được; mất tiện ích bấm chạy.'),
('AI_HANDOFF.md','Bối cảnh và yêu cầu cho trợ lý mới; bạn chủ động yêu cầu đọc.','Web không đổi; AI thiếu quyết định sản phẩm và giới hạn.'),
('SOURCE_MANIFEST.json','Danh sách file nguồn, kích thước/hash và mốc đối chiếu.','Web không đổi; khó chứng minh bộ copy có nguyên vẹn.'),
('THIRD_PARTY_NOTICES.md','Nguồn tham khảo và tình trạng giấy phép.','Web không đổi; mất thông tin xuất xứ cần giữ.')
],[5.2,7.4,4.6])
code('Co_Chill_VSCode_2026_10_17/\n  README.md\n  AI_HANDOFF.md\n  .gitignore\n  SOURCE_MANIFEST.json\n  THIRD_PARTY_NOTICES.md\n  .vscode/tasks.json\n  online-preview/    10 tài nguyên chạy web\n  app/               4 file tham chiếu Python\n  tests/             kiểm thử luật Python\n  tools/             kiểm thử và xuất bài')
p('Các file được copy giữ đúng đường dẫn tương đối để test tìm được module. Không chuyển riêng app.js lên thư mục cha hoặc đổi tên thư mục online-preview mà không sửa các đường dẫn liên quan.')
p('Bộ này không có .git, credential, .openai, database, Word riêng hoặc lịch sử chat. Bạn có thể thử độc lập mà không xóa repo cũ. Bản ZIP là snapshot: sửa thư mục sau khi tạo ZIP không tự cập nhật ZIP.')
p('Không có bước build. Thư mục phục vụ công khai là online-preview; các file test, hướng dẫn nội bộ và nguồn Python tham chiếu không cần xuất bản lên hosting.')

page(12,'Bản Python cũ và phần nào nên lưu riêng')
p('Repo đầy đủ vẫn nằm ở E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai. Các file dưới đây thuộc ứng dụng local riêng ở cổng 8765. Không cần chúng để chạy bộ Site gọn, nhưng hãy lưu repo gốc thay vì xóa để tránh mất một triển khai còn có giá trị.')
table(['Đường dẫn tại repo gốc','Vai trò trong app Python và nếu thiếu'],[
('server.py','Nhận HTTP, phục vụ web/ và API, cookie phiên khách, giới hạn truy cập local. Thiếu: không khởi động app Python.'),
('web/index.html','Khung giao diện Python. Thiếu: không có trang đầu của app Python.'),
('web/style.css','Giao diện Python. Thiếu: trang và bàn cờ mất định dạng.'),
('web/app.js','Frontend gọi /api/ và xử lý thao tác. Thiếu: app Python không tương tác được.'),
('web/favicon.svg','Icon của app Python. Thiếu: chỉ mất icon.'),
('app/service.py','Điều phối API chơi/học, reflection gate phía server, review, dashboard. Thiếu: nghiệp vụ backend lỗi.'),
('app/learning.py','Giải thích mẫu, khuyến nghị và lịch ôn phía Python. Thiếu: luồng học backend lỗi.'),
('app/storage.py','Đọc/ghi SQLite theo phiên khách. Thiếu: backend không lưu và đọc dữ liệu đúng.'),
('app/rules.py và engine.py','Luật và engine Python. Bốn file tham chiếu đã được giữ trong bộ gọn; các file còn lại vẫn ở repo gốc.'),
('app/puzzles.py và __init__.py','Nguồn bài và package Python. Nguồn bài cũng phục vụ xuất fixtures cho Site.'),
('tests/test_learning.py\ntests/test_http.py','Test nghiệp vụ và HTTP của bản Python; không đưa vào bộ gọn vì thiếu backend tương ứng.'),
('Mo_Co_Chill.cmd\ntools/start_local.ps1','Launcher của app Python cổng 8765. Không phải launcher Site cổng 8769.'),
('var/chill.sqlite3','Dữ liệu người học local Python. Xóa là mất tiến độ; sao lưu riêng tư, không đưa vào GitHub.')
],[6.4,10.8])
p('Hai app không tự đồng bộ. Sửa web/app.js không đổi online-preview/app.js. Chạy python server.py --port 8765 mở bản Python; chạy python -m http.server ... --directory online-preview mở bản Site. Backend local chưa phải dịch vụ tài khoản sẵn sàng công khai.')

page(13,'Tài liệu và file khác có thể tách khỏi repo gọn')
p('“Không cần để chạy” không có nghĩa “nên xóa”. Hãy giữ thư mục gốc làm kho lưu trữ; chỉ chuyển bộ gọn sang nơi phát triển mới. Việc chuyển công cụ không yêu cầu dọn hoặc viết lại lịch sử Git.')
table(['Vị trí ở repo gốc','Giữ để làm gì và ảnh hưởng nếu không copy'],[
('Answer V3.docx\nCodex.docx\nXiangqiAI_V3_Final_Pack.zip\ndocs/spec/00…17','Ý định, yêu cầu và đặc tả gốc. Không ảnh hưởng chạy web, nhưng mất căn cứ sản phẩm nếu xóa. Giữ nguyên tài liệu gốc.'),
('docs/UIUX_DECISION_REGISTER_2026_10_15.md\ndocs/RULES_UIUX_2026_10_15.md','Quyết định UI/UX và luật đã đối chiếu. Có ích trước khi sửa luồng hoặc thêm luật.'),
('docs/CO_CHILL_V5_QUESTIONS.md\ndeliverables/Co_Chill_V5_Cap_nhat_Tai_khoan.docx','Bộ V5 46 câu đã gộp tài khoản. Bạn chỉ điền bản Word cập nhật; tài liệu không phải chức năng đã triển khai.'),
('docs/V4_BO_SUNG_UI_UX_KIEN_TRUC_LUAT.md\ndeliverables/Co_Chill_V4_*.docx','Câu hỏi và tài liệu các vòng trước. Không cần runtime; lưu để tra quyết định và phần chưa trả lời.'),
('IMPLEMENTATION_REPORT_CODEX.md\nTEST_REPORT_CODEX.md\nAI_USAGE.md\nKNOWN_LIMITATIONS.md\nPRIVACY.md','Báo cáo triển khai, kiểm thử, AI, giới hạn và dữ liệu. Nên giữ ở kho tài liệu; một số nội dung mô tả bản Python/lịch sử.'),
('BAO_CAO_CHATGPTSITE_V4_2026_10_15.md\nBAO_CAO_UIUX_V5_2026_10_16.md','Báo cáo hai lần nâng bản online. Giúp biết việc đã làm và phần còn thiếu.'),
('PROJECT_MEMORY_VI.md\nPROJECT_INVENTORY_VI.md\ninventory/\ndeliverables/Co_Chill_So_tay_hien_trang.html','Kiểm kê và bộ nhớ dự án. Snapshot cũ không thay thế kiểm tra source hiện tại.'),
('.git/','Lịch sử commit/nhánh/remote. Không cần chạy web; không đưa vào ZIP gọn. Giữ repo gốc hoặc clone đúng remote nếu cần lịch sử.'),
('var/ và artifacts/','Runtime, bản đóng gói, log và ảnh kiểm tra. Không copy hàng loạt: có thể có dữ liệu riêng và trạng thái deploy.'),
('docs/materials/ và tài liệu chủ sở hữu','Nguồn trả lời riêng. Chỉ cung cấp cho trợ lý khi bạn chủ động chọn; không cần phục vụ lên web.')
],[8.4,8.8])
p('Các hướng dẫn và nhật ký khác như GITHUB_CHECKLIST_VI.md, CONTRIBUTING.md, DEVELOPMENT_JOURNAL.md, kế hoạch và báo cáo nghiên cứu cũng không phải runtime. Công cụ tạo Word, render, kiểm kê trong tools/ không cần trong bộ phát triển Site gọn. __pycache__, *.pyc có thể sinh lại; đừng nhầm chúng với mã nguồn.')

page(14,'Muốn đổi một phần sản phẩm thì sửa file nào')
table(['Bạn muốn đổi','Bắt đầu từ đâu','Cần kiểm tra lại'],[
('Màu và bố cục','online-preview/style.css','Desktop, mobile, tương phản và bàn cờ.'),
('Logo hoặc menu chung','logo-co-chill.png, index.html','Đường dẫn ảnh, nhãn, menu điện thoại.'),
('Trang chủ và lời dẫn','app.js: home(), howto(), about()','Link và nút vẫn dẫn đúng bước.'),
('Điều kiện trả lời trước đáp án','learning.js + experience.js + app.js','Chặn đáp án, hỗ trợ, nháp, retry và ghi nhận.'),
('Thêm bài tập','app/puzzles.py + exporter','Nước hợp lệ, nguồn gốc, fixtures xuất khớp.'),
('Luật hoặc sức máy','xiangqi.js; đối chiếu app/rules.py/engine.py','Toàn bộ luật, parity, tìm kiếm và hiệu năng.'),
('Ký phổ và nhập ván','game-tools.js','Round trip ký phổ, tình huống nhập sai, an toàn dữ liệu.'),
('Lịch ôn và bài liên quan','learning.js','Mốc hoàn thành, không trùng, tiếp tục bài dở.'),
('Dashboard học tập','experience.js + app.js: profile()','Số mẫu, lần có hỗ trợ, không thổi phồng kết quả.'),
('Tài khoản và quản trị V5','Thiết kế backend/auth/database mới, rồi nối app.js','Phân quyền server, đồng bộ, xử lý xung đột, export/xóa.')
],[4.3,6.3,6.6])
h('Nhịp sửa dễ kiểm soát')
bullet('Chọn một thay đổi có thể nhìn hoặc kiểm tra được; yêu cầu AI nêu file sẽ sửa.')
bullet('Xem diff trước và sau. Không nhận một lần viết lại toàn bộ framework chỉ để đổi màu hoặc bố cục.')
bullet('Chạy test phù hợp, thử luồng chính trên browser, rồi lưu một commit rõ nội dung.')
bullet('Khi thay cấu trúc dữ liệu lưu, yêu cầu tương thích hoặc chuyển đổi dữ liệu cũ; không xóa localStorage để che lỗi.')
p('Bạn không phải chỉnh tất cả file mỗi lần. Kiến trúc chia phần giúp đổi đúng nơi; nhưng một yêu cầu liên quan nhiều tầng, như tài khoản thật, sẽ cần thay đổi nhiều file và bổ sung thành phần mới.')

page(15,'Kiểm thử GitHub và phát hành')
h('Kiểm thử bộ gọn')
code('python tools/check_all.py\n# Nếu Node trong PATH quá cũ, dùng:\npython tools/check_all.py --node "DUONG_DAN_NODE_MOI/node.exe"')
p('Công cụ kiểm tra không cài package hoặc tự tải code. Node 6.14.0 trên PATH máy hiện tại không đủ; việc này ảnh hưởng test, không ngăn Python phục vụ web. Cài Node hiện đại từ nguồn chính thức là bước môi trường riêng; bộ bàn giao không tự sửa cấu hình máy.')
report_path=ROOT/'var/migration-2026-10-17/package_report.json'
report=json.loads(report_path.read_text(encoding='utf-8'))
assert report['file_count']==29 and report['runtime_hashes_match'] and report['zip_exact_filelist_and_hashes_match']
bullet('Đạt 12/12 nhóm kiểm tra: cú pháp cả sáu JavaScript, ba bộ test Node, luật Python, dữ liệu bài và đối chiếu JS/Python.')
bullet('Đạt 51 test JavaScript và 30 test luật Python. Cả 12 bài xuất khớp nguồn và tập nước hợp lệ.')
bullet('Đối chiếu đạt trên 205 thế, 5.468 thế con, 25 ca FEN và 153 nhánh engine. Đây là phạm vi kiểm thử cụ thể, chưa phải chứng nhận toàn bộ luật giải đấu.')
bullet('Bộ kiểm thử chạy bằng Python 3.12.14 và Node 24.19.0. Runner cũng đã từ chối Node 6.14.0 với thông báo rõ ràng.')
bullet('29 file trong ZIP khớp bản thư mục; 23 file sao chép khớp nguồn. Không có cache Python hoặc dữ liệu người học trong ZIP.')
p('Không có build hay typecheck TypeScript trong dự án hiện tại. Kiểm tra cú pháp JavaScript và các test hành vi là bước phù hợp. Các bài test bộ gọn không thay cho toàn bộ suite của app Python cũ.')
h('Commit push và deploy là ba thao tác khác nhau')
p('Commit tạo mốc trên máy. Push gửi các mốc lên GitHub. Deploy đưa một phiên bản lên hosting. Bộ ZIP chưa có lịch sử Git hoặc remote; mở bằng VS Code không tự tạo kết nối với repo xiangqi07.')
p('Để tiếp tục repo GitHub hiện có, clone repo bằng tài khoản có quyền, tạo nhánh làm việc rồi đưa các thay đổi đã chọn vào. Không lấy ZIP rồi force push đè lịch sử repo cũ. Nếu chỉ muốn đọc/chạy, chưa cần Git.')
code('git status\ngit diff\n# Chọn đúng các file đã kiểm tra để stage\ngit add online-preview README.md\ngit diff --cached\ngit commit -m "Update the reviewed change"')
p('Lệnh trên dành cho thư mục đã là repo Git. Trước push, kiểm tra remote và tên nhánh. Trợ lý có thể làm Git khi được cấp quyền; yêu cầu hỏi bạn trước push, deploy, đổi Public hoặc xóa dữ liệu vẫn phải được giữ.')
h('Đưa lên web công khai sau này')
p('Bản hiện tại có thể phục vụ bằng hosting tĩnh với thư mục xuất bản online-preview, không cần build. VS Code không phải hosting; máy bạn tắt thì local server ngừng. Link chatgpt.site thuộc dự án Sites đã có, không đi theo ZIP. Muốn cập nhật link đó cần quyền và công cụ publish của Sites; chuyển hosting là một cấu hình khác.')
p('Khi V5 có tài khoản, cần triển khai thêm backend/auth/database tương thích. Không coi việc tải 10 file lên Vercel là đã có đăng ký, quản trị hoặc đồng bộ.')

page(16,'Bàn giao cho trợ lý mới và việc cần làm tiếp')
p('Mở AI_HANDOFF.md trong bộ gọn và yêu cầu trợ lý đọc. Không giả định một extension tự đọc mọi tài liệu hoặc tự có lịch sử chat này. Đoạn dưới có thể dùng làm yêu cầu đầu tiên.')
code('Tôi phát triển Cờ Chill. Hãy đọc README.md và\nAI_HANDOFF.md trước. online-preview là source web;\napp/ chỉ là tham chiếu luật/engine/bài trong bộ gọn.\nChạy kiểm thử, báo đúng phần thật và phần chưa làm.\nGiữ suy nghĩ trước đáp án, luật quyết định nước đi,\nkhông coaching trong ván, không bỏ hoặc nới test.\nKhông tự push, deploy, đổi Public, xóa dữ liệu hoặc\nđăng ký dịch vụ khi tôi chưa duyệt hành động cụ thể.\nTrước tiên giải thích code hiện tại và chỉ ra file\nsẽ sửa cho yêu cầu tiếp theo của tôi.')
h('Tài liệu bổ sung cho vòng V5')
bullet('Bản trả lời V4: E:\\ThanhTQ\\App cờ tướng\\10-14\\V4.docx.')
bullet('Bản UI UX và Luật: E:\\ThanhTQ\\App cờ tướng\\10-13\\UI UX + Luat.docx; phần bổ sung bạn đã chỉ ở trang 47.')
bullet('V5 cập nhật cần hoàn thành: E:\\ThanhTQ\\App cờ tướng\\xiangqi-ai\\deliverables\\Co_Chill_V5_Cap_nhat_Tai_khoan.docx. Chỉ điền một bản V5 này; tài liệu bàn giao code không phải bộ câu hỏi mới.')
p('Trước khi AI thay thiết kế hoặc luật, cung cấp phần trả lời mới có liên quan. V5 phải bao gồm tài khoản, dữ liệu tập trung, đồng bộ và quản trị đã chốt; chưa có tính năng đó trong bản sao hôm nay. Chọn dịch vụ và phương án triển khai sau khi chốt ngân sách, quyền truy cập và tiêu chí nghiệm thu.')
h('Điểm cần nhớ khi đọc các báo cáo cũ')
p('Repo chứa nhiều mốc phát triển. Một số README/báo cáo còn mô tả demo hai bài hoặc app Python chưa công khai. Bản Site được kiểm kê ở đây có 12 bài, giao diện mới và đồng hồ. Khi có mâu thuẫn, kiểm tra đúng source và đúng lần phát hành; không trộn kết quả test của hai ứng dụng.')
h('Nguồn để tự học tiếp')
p('[1] VS Code Workspaces: https://code.visualstudio.com/docs/editing/workspaces/workspaces — mở thư mục dự án. VS Code CLI: https://code.visualstudio.com/docs/configure/command-line — mở folder/file bằng code và --new-window.')
p('[2] Python http.server: https://docs.python.org/3/library/http.server.html — chạy server file tĩnh, bind và directory. Lệnh này dành cho thử local, không phải hệ thống hosting sản phẩm.')
p('Căn cứ kiến trúc: 10 file online-preview, server.py, app/, web/, bộ kiểm thử và checkout Sites có commit ghi ở trang 1. Manifest trong bộ gọn ghi file nguồn và SHA256 để bạn đối chiếu bản copy. Không có quyền tái phân phối mới được cấp bởi việc tạo ZIP; giữ thông tin nguồn và quyết định giấy phép của chủ dự án.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(str(OUT))
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
